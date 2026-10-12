# -*- coding: utf-8 -*-
"""엑셀 작업일지 자동 읽기. 파일 수정시각이 바뀐 것만 다시 읽어 parquet 캐시에 합친다.
- 팽이: 팽이일지/**/*팽이 생산 상세일지.xlsx (일별생산일지), *팽이 일일작업일지.xlsx (날짜 시트: 수확·병당중량·QC)
- 새송이: 새송이일지/**/*새송이 일일작업일지.xlsx (날짜 시트: 작업이력·생산kg), *새송이 상세일지.xlsx (규격별 생산)
- 배양실: 배양소일지/**/*배양소관리.xlsx (접종·팽이균긁기·새송이균긁기·권지)
- 수리 기록(v1.23.0): 권석현 폴더의 기계수리.xlsx · 가습기 수리내역.xlsx · 수리 완료 보고서.xls(수리비)
"""
import os, re, glob, json, datetime as dt, warnings
import numpy as np
import pandas as pd
import openpyxl
from . import config as C

warnings.filterwarnings("ignore", module="openpyxl")
STATE = os.path.join(C.CACHE_DIR, "logs_state.json")
DATASETS = ["paengi_daily", "paengi_harvest", "paengi_qc", "ss_daily", "ss_events", "ss_kgcat",
            "ss_room_pack", "ss_room_in", "ss_room_qc",
            "cult_inoc", "cult_scrape_paengi", "cult_scrape_ss", "cult_kwonji",
            "bottling", "materials", "materials_block", "electric", "fuel", "ledger",
            "spawn_tanks", "tank_daily", "loads", "sales", "repairs", "repair_costs"]


def _p(name):
    return os.path.join(C.CACHE_DIR, f"logs_{name}.parquet")


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _ts(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return pd.Timestamp(v).normalize()
    return pd.NaT


# ───────────────────────── 팽이 ─────────────────────────
def parse_paengi_detail(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "일별생산일지" not in wb.sheetnames:
        return {"ledger": parse_ledger(wb, "팽이", path)}
    ws = wb["일별생산일지"]
    data = list(ws.iter_rows(values_only=True))
    cols = {}
    for r in data[:8]:
        for j, v in enumerate(r):
            if not isinstance(v, str):
                continue
            s = v.strip()
            if s == "발이실": cols["balsil"] = j
            elif s == "생육실": cols["saengsil"] = j
            elif "입상" in s and "병" in s: cols["ipsang"] = j
            elif s.replace(" ", "") == "합계(kg)": cols["total_kg"] = j
    rows = []
    for r in data:
        d = _ts(r[0])
        if pd.isna(d):
            continue
        rec = {"date": d}
        for k, j in cols.items():
            rec[k] = _num(r[j]) if j < len(r) else None
        rows.append(rec)
    return {"paengi_daily": pd.DataFrame(rows), "ledger": parse_ledger(wb, "팽이", path)}


def _qc_rows(ws_rows, date):
    out, qrow = [], None
    for i in range(59, min(len(ws_rows), 90)):
        v = ws_rows[i][0] if ws_rows[i] else None
        if v and str(v).strip().startswith("QC"):
            qrow = i
            break
    if qrow is None:
        return out
    head = str(ws_rows[qrow][0]).replace(" ", "")
    if "본수확/점검" in head:
        for i in range(qrow + 3, min(qrow + 11, len(ws_rows))):
            r = ws_rows[i]
            room, day, sc = r[0], r[1], (r[24] if len(r) > 24 else None)
            if isinstance(room, (int, float)) and isinstance(sc, (int, float)):
                out.append({"date": date, "room": int(room), "day_n": _num(day), "score": float(sc),
                            "kind": "본수확", "note": str(r[25])[:60] if len(r) > 25 and r[25] else None, "fmt": "B"})
    else:
        r = ws_rows[qrow]
        for c_room, c_sc, c_rem, kind in ((2, 4, 8, "본수확"), (18, 22, 23, "점검수확")):
            room = r[c_room] if len(r) > c_room else None
            sc = r[c_sc] if len(r) > c_sc else None
            if isinstance(room, (int, float)) and isinstance(sc, (int, float)):
                rem = r[c_rem] if len(r) > c_rem else None
                m = re.search(r"(\d+)일", str(rem or ""))
                out.append({"date": date, "room": int(room), "day_n": int(m.group(1)) if m else None,
                            "score": float(sc), "kind": kind, "note": str(rem)[:60] if rem else None, "fmt": "A"})
    return out


def parse_paengi_work(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    H, Q = [], []
    for sn in wb.sheetnames:
        if not re.fullmatch(r"\d{6}", sn):
            continue
        ws = wb[sn]
        rows = list(ws.iter_rows(min_row=1, max_row=90, max_col=27, values_only=True))
        # 날짜는 시트 이름(YYMMDD)에서 잡는다. 안쪽 날짜칸(Y4)은 전날 시트를 복사해 만들면
        # 그대로 남아 있어서, 어제 날짜를 단 빈 시트가 어제 기록을 덮어써 버린다.
        # (새송이 일지 읽는 쪽도 원래 시트 이름을 쓴다. 같은 방식으로 맞춘 것.)
        d = pd.Timestamp("20" + sn[:2] + "-" + sn[2:4] + "-" + sn[4:6])
        if pd.isna(d):
            continue
        g = lambda r, c: rows[r - 1][c - 1] if len(rows) >= r and len(rows[r - 1]) >= c else None
        H.append({"date": d, "harvest_rooms": g(63, 2), "bottles": _num(g(65, 2)),
                  "g_per_bottle": _num(g(63, 24)), "grade_top": _num(g(64, 21)),
                  # 수확 병수(B64)·베이비 병수(D64)·그날 작업 박스(L63)·누적 박스(L64) — 배치 보고서 수확 표에 쓴다
                  "harvest_bottles": _num(g(64, 2)), "baby_bottles": _num(g(64, 4)),
                  "boxes": _num(g(63, 12)), "boxes_cum": _num(g(64, 12))})
        Q += _qc_rows(rows, d)
    return {"paengi_harvest": pd.DataFrame(H), "paengi_qc": pd.DataFrame(Q)}


# ───────────────────────── 새송이 ─────────────────────────
def ss_rooms(s):
    """'2-3,4,5' / '1-10 / 2-1' / '2-56'(오기) → ['B03','B04','B05']"""
    if s is None:
        return []
    out = []
    for part in re.split(r"[/;]", str(s)):
        m = re.match(r"\s*([1-4])\s*-\s*([\d,\s.]+)", part)
        if not m:
            continue
        blk = C.SS_BLOCK[m.group(1)]
        mx = 14 if blk == "D" else 10
        for n in re.findall(r"\d+", m.group(2)):
            n = int(n)
            if n > mx:
                out += [f"{blk}{int(ch):02d}" for ch in str(n) if ch != "0"]
            else:
                out.append(f"{blk}{n:02d}")
    return out


LBL = {"수확": "수확", "솎기": "솎기", "입상/수량": "입상", "입상": "입상", "뒤집기": "뒤집기", "폐상": "폐상", "탈병": "탈병"}


def parse_ss_work(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows, events = [], []
    for sn in wb.sheetnames:
        if not re.fullmatch(r"\d{6}", sn):
            continue
        d = pd.Timestamp("20" + sn[:2] + "-" + sn[2:4] + "-" + sn[4:6])
        grid = list(wb[sn].iter_rows(min_row=30, max_row=42, values_only=True))
        rec = {"date": d, "kg": None, "bottles": None}
        for r in grid:
            if len(r) > 4 and r[1] == "Total" and rec["kg"] is None:
                rec["kg"] = _num(r[4])
            lbl = r[20] if len(r) > 20 else None
            if isinstance(lbl, str) and lbl.strip() in LBL:
                k = LBL[lbl.strip()]
                raw = r[22] if len(r) > 22 else None
                raw = None if raw in (0, "0", "") else raw
                rec[k] = str(raw) if raw is not None else None
                if k == "입상":
                    rec["bottles"] = _num(r[24]) if len(r) > 24 else None
                for rm in ss_rooms(raw):
                    events.append({"date": d, "event": k, "room": rm, "raw": str(raw)})
        rows.append(rec)
    return {"ss_daily": pd.DataFrame(rows), "ss_events": pd.DataFrame(events)}


def parse_ss_detail(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "일별작업일지" not in wb.sheetnames:
        return {"ledger": parse_ledger(wb, "새송이", path)}
    data = list(wb["일별작업일지"].iter_rows(values_only=True))
    cat, kgcol = None, None
    for r in data[:8]:
        if r and r[1] in ("손질", "깎기", "뿌리", "미니"):
            cat = r
        for j, v in enumerate(r):
            if isinstance(v, str) and v.replace(" ", "").startswith("합계(kg)"):
                kgcol = j
    if cat is None:
        return {"ledger": parse_ledger(wb, "새송이", path)}
    rows = []
    for r in data:
        d = _ts(r[0])
        if pd.isna(d):
            continue
        rec = {"date": d, "kg": _num(r[kgcol]) if kgcol is not None else None}
        for j in range(1, min(14, len(r))):
            c = cat[j] if j < len(cat) else None
            if isinstance(c, str):
                rec[c] = rec.get(c, 0) + (r[j] if isinstance(r[j], (int, float)) else 0)
        rows.append(rec)
    return {"ss_kgcat": pd.DataFrame(rows), "ledger": parse_ledger(wb, "새송이", path)}


# ───────────────────────── 배양실 ─────────────────────────
_BLANK = ("", "nan", "none", "nat", "-")


def _coalesce(df, *names):
    """월 블록마다 헤더 이름이 다른 열(예: '비고/발이실' vs '비고')을 앞 이름 우선으로 합친다."""
    out = None
    for n in names:
        if n not in df:
            continue
        col = df[n]
        if out is None:
            out = col.copy(); continue
        have = out.notna() & ~out.astype(str).str.strip().str.lower().isin(_BLANK)
        out = out.where(have, col)
    return out


def _blocks(ws, width):
    """'일자' 헤더가 있는 모든 블록(가로로 월별, 세로로 연도별)을 읽는다."""
    data = list(ws.iter_rows(values_only=True))
    out = []
    for i, r in enumerate(data):
        for j, v in enumerate(r or ()):
            if v != "일자":
                continue
            cols = [data[i][j + k] if j + k < len(r) else None for k in range(width)]
            for rr in data[i + 1:]:
                d = rr[j] if j < len(rr) else None
                if d == "일자":
                    break
                if isinstance(d, (dt.datetime, dt.date)):
                    rec = {str(cols[k]): (rr[j + k] if j + k < len(rr) else None) for k in range(width)}
                    rec["일자"] = pd.Timestamp(d).normalize()
                    out.append(rec)
    return pd.DataFrame(out)


def _first_date(v, ref):
    """'1/16,19' 같은 다중 표기 → 첫 날짜(연도는 기준일 기준으로 추정)."""
    if v is None or v is pd.NaT or (isinstance(v, float) and v != v):
        return pd.NaT
    if isinstance(v, (dt.datetime, dt.date, pd.Timestamp)):
        if pd.isna(v):
            return pd.NaT
        t = pd.Timestamp(v).normalize()
        if t > ref + pd.Timedelta(days=3):          # 연도 오기(예: 2026-11-18 → 2025-11-18)
            t = t.replace(year=t.year - 1)
        return t
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})", str(v or ""))
    if m:
        mo, d = int(m.group(1)), int(m.group(2))
        y = ref.year if mo <= ref.month else ref.year - 1
        try:
            return pd.Timestamp(year=y, month=mo, day=d)
        except ValueError:
            return pd.NaT
    return pd.NaT


def parse_cult(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    num = lambda s: pd.to_numeric(s, errors="coerce")

    def sheet(cands):
        for n in wb.sheetnames:
            if any(c in n for c in cands):
                return wb[n]
        return None

    ws = sheet(["접종수량", "접종"])
    if ws is not None:
        J = _blocks(ws, 6)
        if not J.empty and "품목" in J:
            J = J[J["품목"].isin(["팽이", "새송이"])].copy()
            J["접종량"], J["손실"] = num(J.get("접종량")), num(J.get("손실"))
            out["cult_inoc"] = J[["일자", "품목", "접종량", "손실"]].dropna(subset=["접종량"])
    for cands, key, roomcol in ((["팽이균긁기", "팽이긁기"], "cult_scrape_paengi", "발이실"),
                                (["새송이균긁기", "새송이긁기"], "cult_scrape_ss", "생육실")):
        ws = sheet(cands)
        if ws is None:
            continue
        P = _blocks(ws, 10)
        if P.empty:
            continue
        P["배양출고"], P["불량"] = num(P.get("배양출고")), num(P.get("불량"))
        P["균긁기량"] = num(_coalesce(P, "균긁기", "균긁기량"))
        _inoc = _coalesce(P, "접종일", "접종일자")
        P["접종일"] = [_first_date(v, d) for v, d in zip(_inoc if _inoc is not None else [None] * len(P), P["일자"])]
        _room = _coalesce(P, roomcol, roomcol + "NO", "비고/" + roomcol)
        P["방"] = _room.astype(str).str.extract(r"(\d+)")[0] if _room is not None else None
        P["방_raw"] = _room.astype(str) if _room is not None else None
        P["탱크"] = P["탱크 NO."].astype(str) if "탱크 NO." in P else None
        P = P[P["배양출고"] > 0]
        out[key] = P[["일자", "접종일", "배양출고", "불량", "균긁기량", "방", "방_raw", "탱크"]].rename(columns={"일자": "긁기일"})
    ws = sheet(["권지"])
    if ws is not None:
        K = _blocks(ws, 10)
        if not K.empty:
            K["입상량"] = num(_coalesce(K, "입상량", "균긁기량", "작업수량"))
            K["불량"], K["권지작업량"] = num(K.get("불량")), num(K.get("권지작업량"))
            _ino = _coalesce(K, "접종일자", "접종일")
            K["접종일"] = [_first_date(v, d) for v, d in zip(_ino if _ino is not None else [None] * len(K), K["일자"])]
            ip = _coalesce(K, "입상일자", "긁기 일자", "긁기일자", "균긁기일")
            K["입상일"] = [_first_date(v, d) for v, d in zip(ip if ip is not None else [None] * len(K), K["일자"])]
            bal = _coalesce(K, "발이실", "비고/발이실", "비고")     # 1~2월 블록은 '비고' 열에 발이실 번호가 적힌다
            b2 = pd.to_numeric(bal.astype(str).str.extract(r"(\d+)")[0], errors="coerce") if bal is not None else None
            K["발이실"] = b2.where(b2.between(1, 18)) if b2 is not None else None
            K["생육실"] = num(K.get("생육실"))
            K = K[K["권지작업량"] > 0]
            out["cult_kwonji"] = K[["일자", "접종일", "입상일", "발이실", "생육실", "입상량", "불량", "권지작업량"]].rename(columns={"일자": "권지일"})
    return out


def parse_ss_room_pack(path):
    """'새송이 방별 포장량.xlsx' — 포장량 시트(포장일·포장 방·미니/새송이 kg·박스수), 방별 시트(입상일·방·입상병수).
    방 단위 수확량이 있는 유일한 자료(2025-10~2026-04, 관리자 교체 전)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    P, I = [], []
    if "포장량" in wb.sheetnames:
        hdr = None
        for r in wb["포장량"].iter_rows(min_row=1, max_row=6, values_only=True):
            if r[1] == "포장 일" and r[2] == "포장 방":
                hdr = [str(c).strip() if c is not None else "" for c in r]
                break
        for r in wb["포장량"].iter_rows(min_row=4, values_only=True):
            d = r[1]
            if not hasattr(d, "year") or r[2] is None:
                continue
            rooms = ss_rooms(r[2])
            if len(rooms) != 1:
                continue
            rec = {"date": pd.Timestamp(d).normalize(), "room": rooms[0], "방_raw": str(r[2]).strip(),
                   "미니kg": _num(r[3]), "새송이kg": _num(r[4]), "kg": _num(r[5])}
            if hdr:
                for j, h in enumerate(hdr):
                    if re.fullmatch(r"[\d.]+kg", h) and j >= 7:
                        rec["box_" + h] = _num(r[j])
            P.append(rec)
    if "방별" in wb.sheetnames:
        for r in wb["방별"].iter_rows(min_row=3, values_only=True):
            d = r[1]
            if not hasattr(d, "year") or r[2] is None:
                continue
            rooms = ss_rooms(r[2])
            if len(rooms) == 1 and _num(r[3]):
                I.append({"입상일": pd.Timestamp(d).normalize(), "room": rooms[0], "입상병수": _num(r[3])})
    wb.close()
    return {"ss_room_pack": pd.DataFrame(P), "ss_room_in": pd.DataFrame(I)}


SS_QC_ITEMS = [("생육", "생육"), ("대모양", "대모양"), ("갓상태", "갓상태"), ("병해", "병해"), ("생산량", "생산량점")]


def parse_ss_room_qc(path):
    """정규철 '생산량 (M월-M월).xlsx' 의 '실별QC' 시트 → 방 단위 새송이 배치 결과.
    입상일·재배실·솎기·포장종료·입상병수·수확kg·병당g 와 QC 항목 점수(생육·대모양·갓상태·병해·생산량, 각 10점)·종합·비고.
    '새송이 방별 포장량.xlsx'(2025-10~2026-04)이 끊긴 2026-05 이후 구간을 메우는 유일한 방 단위 수확 자료."""
    try:
        d = pd.read_excel(path, sheet_name="실별QC", header=2)
    except Exception:
        return {}
    d.columns = [str(c).strip() for c in d.columns]
    if "재배실" not in d or "입상일" not in d:
        return {}
    d = d[d["재배실"].notna()].copy()
    d["입상일"] = pd.to_datetime(d["입상일"], errors="coerce")
    d = d[d["입상일"].notna()]
    out = []
    for _, r in d.iterrows():
        rooms = ss_rooms(r["재배실"])
        if len(rooms) != 1:
            continue
        rec = {"입상일": r["입상일"].normalize(), "room": rooms[0], "방_raw": str(r["재배실"]).strip(),
               "솎기": pd.to_datetime(r.get("솎기"), errors="coerce"),
               "포장종료": pd.to_datetime(r.get("포장종료"), errors="coerce"),
               "입상병수": _num(r.get("입상량")), "수확kg": _num(r.get("수확량")), "병당g": _num(r.get("평균중량")),
               "QC종합": _num(r.get("종합")),
               "발이메모": str(r.get("발이 상태")).strip() if pd.notna(r.get("발이 상태")) else None,
               "비고": str(r.get("비고")).strip() if pd.notna(r.get("비고")) else None}
        for src, dst in SS_QC_ITEMS:
            rec[dst] = _num(r.get(src))
        if rec["수확kg"] and rec["입상병수"] and not rec["병당g"]:
            rec["병당g"] = 1000 * rec["수확kg"] / rec["입상병수"]
        out.append(rec)
    return {"ss_room_qc": pd.DataFrame(out)}


def parse_ledger(wb, crop, path):
    """생산 상세일지의 '입출고내역' 시트 → 입출고 원장 (재고/입고/출고 행). 파일명의 연·월을 month 로."""
    if "입출고내역" not in wb.sheetnames:
        return pd.DataFrame()
    m = re.search(r"(\d{4})년\s*(\d{1,2})월", os.path.basename(path))
    month = f"{m.group(1)}-{int(m.group(2)):02d}" if m else None
    hdr, rows, seq = None, [], 0
    for r in wb["입출고내역"].iter_rows(min_row=1, values_only=True):
        if r and r[0] == "날짜":
            hdr = [str(c).replace("\n", "").replace(" ", "") if c else "" for c in r]
            continue
        if hdr is None or not r or not isinstance(r[0], (dt.datetime, dt.date)):
            continue
        d = dict(zip(hdr, r))
        typ = d.get("입출고")
        if typ not in ("재고", "입고", "출고", "클레임출고", "기타출고", "폐기", "매입", "실사"):
            continue
        seq += 1
        rows.append({"crop": crop, "month": month or pd.Timestamp(r[0]).strftime("%Y-%m"), "seq": seq, "date": _ts(r[0]), "type": typ,
                     "cust": d.get("거래처"), "sub": d.get("거래처세부"), "gubun": d.get("구분"), "eco": d.get("친환경"), "grade": d.get("등급"),
                     "name": d.get("규격명"), "spec": d.get("규격"), "unit": _fnum(d.get("단위중량")), "box": _fnum(d.get("수량(box)")),
                     "kg": _fnum(d.get("중량(kg)")), "price": _fnum(d.get("단가(원)")), "amt": _fnum(d.get("금액(원)")), "note": d.get("비고")})
    return pd.DataFrame(rows)


# ───────────────────────── 원가 자료 (권석현 폴더) ─────────────────────────
def _fnum(v):
    """숫자 또는 '조은410,420' 같은 문자열에서 첫 숫자."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        m = re.search(r"\d[\d,]*\.?\d*", v)
        if m:
            try:
                return float(m.group(0).replace(",", ""))
            except ValueError:
                return None
    return None


def parse_bottling(path):
    """팽이재배일보 / 새송이작업일보 '종합관리' 시트 → 날짜별 입병 병수(혼합기별 수량 합), 살균 수량, 입병 회수."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "종합관리" not in wb.sheetnames:
        return {}
    rows = list(wb["종합관리"].iter_rows(values_only=True))
    hdr_i = next((i for i, r in enumerate(rows[:10]) if r and r[0] == "입병일"), None)
    if hdr_i is None:
        return {}
    hdr = [str(c).replace("\n", "").strip() if c else "" for c in rows[hdr_i]]
    ci = next((i for i, h in enumerate(hdr) if h.startswith("혼합기별")), None)
    cs = next((i for i, h in enumerate(hdr) if h.startswith("살균수량")), None)
    cp = hdr.index("품목") if "품목" in hdr else 1
    crop_default = "팽이" if "팽이" in os.path.basename(path) else "새송이"
    agg = {}
    for r in rows[hdr_i + 1:]:
        d = _ts(r[0] if r else None)
        if pd.isna(d):
            continue
        q = r[ci] if ci is not None and ci < len(r) else None
        if not isinstance(q, (int, float)) or q <= 0:
            continue
        crop = r[cp] if cp < len(r) and isinstance(r[cp], str) and r[cp].strip() in ("팽이", "새송이") else crop_default
        k = (d, crop)
        a = agg.setdefault(k, {"date": d, "crop": crop, "bottles": 0.0, "sterilized": 0.0, "batches": 0})
        a["bottles"] += q
        sv = r[cs] if cs is not None and cs < len(r) else None
        a["sterilized"] += sv if isinstance(sv, (int, float)) else 0
        a["batches"] += 1
    return {"bottling": pd.DataFrame(list(agg.values()))}


def parse_materials(path):
    """원재료입고.xlsx (연도 시트, 월 블록) → 월·품목별 전월재고/입고량/사용량/현재고/단가 + 작물별 사용kg·금액, 블록의 입병수·입병일수."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    M, B = [], []
    for sn in wb.sheetnames:
        if not re.fullmatch(r"\d{4}", sn):
            continue
        rows = list(wb[sn].iter_rows(values_only=True))
        heads = [i for i, r in enumerate(rows) if r and isinstance(r[0], str) and "원재료입고" in r[0]]
        for bi, i in enumerate(heads):
            end = heads[bi + 1] if bi + 1 < len(heads) else len(rows)
            m = re.search(r"(\d{4})년\s*(\d{1,2})월", rows[i][0])
            if not m:
                continue
            month = f"{m.group(1)}-{int(m.group(2)):02d}"
            hi = next((j for j in range(i, min(i + 5, end)) if rows[j] and any(isinstance(c, str) and "입고량" in c for c in rows[j])), None)
            if hi is None:
                continue
            hdr = rows[hi]
            col = {}
            for j, c in enumerate(hdr):
                if not isinstance(c, str):
                    continue
                t = c.replace(" ", "")
                if t in ("전월제고", "전월재고"): col["prev"] = j
                elif t == "입고량": col["in"] = j
                elif t == "사용량": col["use"] = j
                elif t in ("현제고", "현재고"): col["stock"] = j
                elif t == "단가": col["price"] = j
                elif t == "팽이": col.setdefault("p_won", j) if j < 17 else col.setdefault("p_kg", j)
                elif t == "새송이": col.setdefault("s_won", j) if j < 17 else col.setdefault("s_kg", j)
            # 블록 머리(헤더 위 두 줄)의 입병일수·입병수
            info = {"month": month}
            for j in range(max(i, hi - 3), hi):
                r = rows[j]
                if not r:
                    continue
                for pk, sk in (("p_won", "s_won"), ("p_kg", "s_kg")):
                    if pk in col and col[pk] < len(r) and isinstance(r[col[pk]], (int, float)):
                        v_p = r[col[pk]]
                        v_s = r[col[sk]] if sk in col and col[sk] < len(r) and isinstance(r[col[sk]], (int, float)) else None
                        if v_p > 1000:
                            info.setdefault("팽이_입병수", v_p); info.setdefault("새송이_입병수", v_s)
                        elif v_p > 0:
                            info.setdefault("팽이_입병일수", v_p); info.setdefault("새송이_입병일수", v_s)
            B.append(info)
            j = hi + 1
            while j < end - 1:
                r, r2 = rows[j], rows[j + 1]
                if r and isinstance(r[0], str) and r[0].strip() and r2 and r2[1] == "수량":
                    item = r[0].strip().replace(" ", "")
                    rec = {"month": month, "item": item}
                    for k, key in (("prev", "전월재고"), ("in", "입고량"), ("use", "사용량"), ("stock", "현재고"), ("price", "단가")):
                        rec[key] = _fnum(r2[col[k]]) if k in col and col[k] < len(r2) else None
                    if rec["단가"] is not None and not (50 <= rec["단가"] <= 20000):      # '조은410,420' 같은 표기 오독 방지
                        m2 = re.findall(r"\d{3,5}", str(r2[col["price"]])) if "price" in col else []
                        rec["단가"] = float(m2[0]) if m2 and 50 <= float(m2[0]) <= 20000 else None
                    rec["팽이_kg"] = _fnum(r[col["p_kg"]]) if "p_kg" in col and col["p_kg"] < len(r) else None
                    rec["새송이_kg"] = _fnum(r[col["s_kg"]]) if "s_kg" in col and col["s_kg"] < len(r) else None
                    rec["팽이_원"] = _fnum(r2[col["p_won"]]) if "p_won" in col and col["p_won"] < len(r2) else None
                    rec["새송이_원"] = _fnum(r2[col["s_won"]]) if "s_won" in col and col["s_won"] < len(r2) else None
                    M.append(rec)
                    j += 2
                else:
                    j += 1
    return {"materials": pd.DataFrame(M), "materials_block": pd.DataFrame(B)}


def parse_electric(path):
    """전기세.xlsx (YYYY전기세 시트) → 월·계량기별 사용량(kWh)·요금."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for sn in wb.sheetnames:
        m = re.search(r"(\d{4})", sn)
        if not m or "전기" not in sn:
            continue
        year = m.group(1)
        rows = list(wb[sn].iter_rows(values_only=True))
        name_row = next((r for r in rows[:6] if r and any(isinstance(c, str) and "고압" in c for c in r)), None)
        if name_row is None:
            continue
        meters = [(j, c.strip()) for j, c in enumerate(name_row) if isinstance(c, str) and c.strip()]
        seen = set()
        for r in rows:
            if not r or not isinstance(r[0], str):
                continue
            mm = re.fullmatch(r"(\d{1,2})월", r[0].strip())
            if not mm:
                continue
            month = f"{year}-{int(mm.group(1)):02d}"
            if month in seen:            # 아래쪽 누계·보조 표는 무시 (첫 표만)
                continue
            seen.add(month)
            for j, name in meters:
                kwh = r[j] if j < len(r) else None
                cost = r[j + 1] if j + 1 < len(r) else None
                if isinstance(kwh, (int, float)) or isinstance(cost, (int, float)):
                    out.append({"month": month, "meter": name, "kwh": kwh if isinstance(kwh, (int, float)) else None,
                                "cost": cost if isinstance(cost, (int, float)) else None})
    return {"electric": pd.DataFrame(out)}


def parse_fuel(path):
    """'가스 및 경유사용량' → 가스/경유 충진 일자별 수량·단가·금액."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for sn, fuel in (("가스 충진량", "LPG"), ("경유 충진량", "경유")):
        if sn not in wb.sheetnames:
            continue
        for r in wb[sn].iter_rows(min_row=1, values_only=True):
            d = _ts(r[0] if r else None)
            if pd.isna(d):
                continue
            q, u, a = (_fnum(r[1]) if len(r) > 1 else None), (_fnum(r[2]) if len(r) > 2 else None), (_fnum(r[3]) if len(r) > 3 else None)
            if q or a:
                out.append({"date": d, "fuel": fuel, "qty": q, "unit_price": u, "amount": a})
    return {"fuel": pd.DataFrame(out)}


# ───────────────────────── 종균·액탱크 ─────────────────────────
_CODE_RE = re.compile(r"([A-Z]+)(\d+)\s*/\s*([TFP])-(\d{6})-(\d+)")


def _spawn_code(s):
    """'PF24/T-260101-1,NO.10' → dict(계통 PF, 세대 24, 종류 T, 배치 T-260101-1, 탱크NO 10)."""
    if not isinstance(s, str):
        return None
    m = _CODE_RE.search(s)
    if not m:
        return None
    fam, gen, kind, d6, seq = m.groups()
    no = re.search(r"NO\.?\s*(\d+)", s, re.I)
    return {"계통": fam, "세대": int(gen), "종류": kind, "배치": f"{kind}-{d6}-{seq}", "코드": f"{fam}{gen}/{kind}-{d6}-{seq}",
            "탱크NO": int(no.group(1)) if no else None}


def parse_spawn_ledger(path):
    """종균 관리 대장.xlsx (연도 시트): 계대 이력 중 T-(액탱크) 행을 탱크 배치로 추출.
    열: 일시, 품종, 모균, 계대 이력×5, 비고(계통 유래), 입병일, 접종일, 특이사항(교반), 탱크PH, 중량"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        if not re.fullmatch(r"20\d{2}", str(ws.title).strip()):
            continue
        for r in ws.iter_rows(min_row=4, values_only=True):
            if len(r) < 15 or not isinstance(r[1], (dt.datetime, dt.date)):
                continue
            for cell in r[4:9]:
                c = _spawn_code(cell)
                if not c or c["종류"] != "T":
                    continue
                mo = _spawn_code(r[3])
                rows.append({"일시": _ts(r[1]), "품종": r[2], "탱크NO": c["탱크NO"], "코드": c["코드"], "배치": c["배치"], "계통": c["계통"], "세대": c["세대"],
                             "모균": r[3], "모균종류": {"F": "플라스크", "P": "플레이트", "T": "탱크"}.get(mo["종류"]) if mo else None,
                             "유래": r[9], "입병일": _ts(r[10]), "접종일": _ts(r[11]), "교반": r[12],
                             "pH": _fnum(r[13]), "중량": _fnum(r[14])})
    df = pd.DataFrame(rows)
    if not df.empty:
        df["액배양일"] = (df["접종일"] - df["일시"]).dt.days
    return {"spawn_tanks": df}


def parse_cult_worklog(path):
    """배양실 작업 일지.xlsx '작업일지': 주 단위 달력(하루 = 4열×13행). 하루치 → 접종 탱크·pH·입병 배지·새 탱크·비고."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "작업일지" not in wb.sheetnames:
        return {}
    rows = list(wb["작업일지"].iter_rows(values_only=True))
    out = []
    W = max(len(r) for r in rows) if rows else 0
    for i, r in enumerate(rows):
        for b in range(1, W - 3, 4):
            if not isinstance(r[b], (dt.datetime, dt.date)) or i + 12 >= len(rows) or rows[i + 1][b] != "액탱크NO":
                continue
            g = lambda k, j: rows[i + k][b + j] if b + j < len(rows[i + k]) else None
            def txt(v):
                if isinstance(v, (dt.datetime, dt.date)):
                    return f"{v.month}/{v.day}"
                return str(v).strip() if v is not None else ""
            new = [(g(k, 0), g(k, 1)) for k in (9, 10) if g(k, 1)]
            flask = [(g(k, 0), g(k, 1), g(k, 3)) for k in (6, 7) if g(k, 1)]
            rec = {"date": _ts(r[b]), "접종탱크NO": _fnum(g(1, 1)), "종균pH": _fnum(g(1, 3)), "접종탱크코드": txt(g(2, 1)) or None, "계통메모": txt(g(2, 3)) or None,
                   "접종배지": " ".join(x for x in (txt(g(3, 0)), txt(g(3, 1))) if x) or None, "살균후pH": _fnum(g(3, 3)),
                   "입병": " ".join(x for x in (txt(g(4, 0)), txt(g(4, 1))) if x) or None, "살균전pH": _fnum(g(4, 3)),
                   "삼각접종원": "; ".join(f"{txt(a)} {txt(c)}".strip() for a, c, _ in flask) or None,
                   "신규탱크": "; ".join(f"{txt(a)}:{txt(c)}" for a, c in new) or None,
                   "비고": txt(g(11, 1)) or None}
            if any(v not in (None, "") for k, v in rec.items() if k != "date"):
                out.append(rec)
    return {"tank_daily": pd.DataFrame(out)}


# ───────────────────────── 재배 관리자 종합관리 (솥 단위 입병 기록) ─────────────────────────
def _rng(v):
    """'6.2-6.4' → (6.2, 6.4, 6.3). 숫자면 그대로. 못 읽으면 (nan, nan, nan)."""
    if v is None:
        return (np.nan, np.nan, np.nan)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (float(v), float(v), float(v))
    nums = re.findall(r"\d+(?:\.\d+)?", str(v))
    if not nums:
        return (np.nan, np.nan, np.nan)
    a = [float(x) for x in nums[:2]]
    return (min(a), max(a), sum(a) / len(a))


def _tm(v):
    if isinstance(v, dt.time):
        return v.hour * 60 + v.minute
    if isinstance(v, dt.datetime):
        return v.hour * 60 + v.minute
    return np.nan


def _hold(v):
    """'105.0/60' → (105.0, 60)."""
    if v is None:
        return (np.nan, np.nan)
    m = re.findall(r"\d+(?:\.\d+)?", str(v))
    return (float(m[0]) if m else np.nan, float(m[1]) if len(m) > 1 else np.nan)


def parse_master(path):
    """정규철 '재배관리/일지' 통합 파일 → loads(솥 단위 입병 기록: 배지 상태·살균기·접종·균긁기·수확) + sales(월별 매출·출고)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    crop = "팽이" if "팽이" in os.path.basename(path) else "새송이"
    if "종합관리" in wb.sheetnames:
        rows = list(wb["종합관리"].iter_rows(values_only=True))
        hdr = [str(c).replace("\n", "").replace(" ", "") if c else "" for c in rows[2]]
        def col(name, nth=0):
            idx = [i for i, h in enumerate(hdr) if h == name]
            return idx[nth] if len(idx) > nth else None
        recs = []
        for r in rows[3:]:
            if not isinstance(r[0], (dt.datetime, dt.date)) or not r[1]:
                continue
            g = lambda name, nth=0: (r[col(name, nth)] if col(name, nth) is not None and col(name, nth) < len(r) else None)
            ph = _rng(g("pH")); wet = _rng(g("수분")); wt = _rng(g("중량/상자")); fill = _rng(g("충진시간"))
            h1 = _hold(g("1차")); h2 = _hold(g("2차")); it = _rng(g("접종온도"))
            rec = {"crop": str(g("품목") or crop).strip(), "입병일": _ts(r[0]), "혼합기": _fnum(g("혼합기")), "살균기NO": str(g("살균기NO")).strip() if g("살균기NO") is not None else None,
                   "입병시작": _tm(g("입병시작")), "입병종료": _tm(g("입병종료")), "살균수량": _fnum(g("살균수량")) or _fnum(g("혼합기별수량")),
                   "pH": ph[2], "pH_min": ph[0], "pH_max": ph[1], "수분": wet[2], "중량상자": wt[2], "충진시간": fill[2], "배지품온": _fnum(g("배지품온")),
                   "함량": _fnum(g("영야원함량")), "설정1차온도": h1[0], "설정1차분": h1[1], "설정2차온도": h2[0], "설정2차분": h2[1],
                   "접종일": _ts(g("접종일")), "종균": str(g("종균")).strip() if g("종균") else None, "접종온도": it[2], "접종LOSS": _fnum(g("LOSS", 0)), "접종수량": _fnum(g("접종수량")),
                   "균긁기일": _ts(g("균긁기일")), "배양일수": _fnum(g("배양일수")), "발이실": str(g("발이실")).strip() if g("발이실") is not None else None,
                   "긁기불량": _fnum(g("불량", 0)), "입상수량": _fnum(g("입상수량", 0)), "긁기비고": str(g("비고", 2)).strip() if g("비고", 2) else None}
            if crop == "팽이":
                rec.update({"권지일": _ts(g("권지일")), "생육실": _fnum(g("생육실")), "권지불량": _fnum(g("불량", 1)), "권지수량": _fnum(g("입상수량", 1)),
                            "상태": str(g("상태")).strip() if g("상태") else None, "수확일": _ts(g("수확일")), "생육기간": _fnum(g("생육기간")), "평균중량": _fnum(g("평균중량"))})
            else:
                rec.update({"솎기일": _ts(g("솎기일")), "수확일": _ts(g("수확일")), "생육기간": _fnum(g("생육기간"))})
            recs.append(rec)
        L = pd.DataFrame(recs)
        if not L.empty:
            L["입병순번"] = L.groupby("입병일").cumcount() + 1
            L["긁기불량률%"] = 100 * L["긁기불량"] / L["살균수량"]
        out["loads"] = L
    # 매출분석 (연도×월 표)
    sh = "팽이(매출)분석" if "팽이(매출)분석" in wb.sheetnames else ("매출분석" if "매출분석" in wb.sheetnames else None)
    if sh:
        rows = list(wb[sh].iter_rows(values_only=True))
        off = 1 if rows[2][0] is None else 0                    # 새송이 시트는 A열이 비어 있음
        years = []
        for y in rows[2][off + 1:off + 6]:
            m_ = re.match(r"(20\d\d)", str(y or ""))
            if not m_ or m_.group(1) in years:      # 다음 블록의 연도가 다시 시작되면 멈춤
                break
            years.append(m_.group(1))
        blocks = {}
        for i, c in enumerate(rows[1]):
            if c and str(c).strip():
                blocks[str(c).replace(" ", "")] = i
        recs = []
        for r in rows[3:]:
            m = re.fullmatch(r"(\d{1,2})월", str(r[off] or "").strip())
            if not m:
                continue
            for k, y in enumerate(years):
                rec = {"crop": crop, "month": f"{y}-{int(m.group(1)):02d}"}
                for label, key in (("출고량(kg)", "출고kg"), ("금액(천원)", "매출천원"), ("매출(원)", "매출원"), ("평균단가(원/박스)", "단가박스"), ("평균단가(원)/kg", "단가kg"),
                                   ("수확병수(병)", "수확병수"), ("출고량대비중량(g/병)", "병당g")):
                    if label in blocks:
                        v = r[blocks[label] + k]
                        rec[key] = float(v) if isinstance(v, (int, float)) else np.nan
                if "매출천원" in rec:
                    rec["매출원"] = rec.pop("매출천원") * 1000
                recs.append(rec)
        S = pd.DataFrame(recs)
        if not S.empty:
            S = S.dropna(subset=["출고kg", "매출원"], how="all")
            S["단가kg"] = S.get("단가kg", np.nan)
            S.loc[S["단가kg"].isna() & S["출고kg"].gt(0), "단가kg"] = S["매출원"] / S["출고kg"]
        out["sales"] = S
    return out


# ───────────────────────── 수리 기록 (권석현 폴더: 기계수리.xlsx · 가습기 수리내역.xlsx) ─────────────────────────
# 기계수리.xlsx 는 구역별 시트(입병실·접종실·발이실·생육실 …)에 '날짜 | 수리내역 | 사양' 으로 적는다.
# '일백' 시트는 냉동 업체(일백)가 한 공사를 모은 것이라 구역 시트와 겹치는 줄이 많다 → 같은 날 비슷한 내용은 한 줄로 합친다.
# 내용에서 방·장비를 뽑아 '장소'(발이실06, 생육실03, C05, 냉각1실, 살균솥 2호기, 배양실 6번 냉동기 …)와
# 환경 기록의 방 이름('방': 발이실06, 냉각1실, 배양실1 … 쉼표로 여러 개)을 붙인다.
REPAIR_AREA = {"입병실": "입병실", "접종실": "접종·냉각실", "균긁기실": "균긁기실", "새송이": "새송이 재배사", "탈병실": "탈병실",
               "팽이배양실": "팽이 배양실", "세송이배양실": "새송이 배양실", "새송이배양실": "새송이 배양실", "발이실": "팽이 발이실",
               "기타장비": "기타 장비", "권지실": "권지실", "생육실": "팽이 생육실", "포장실": "포장·저장고", "콤푸레샤": "콤프레샤",
               "일백": None}
_AREA_BY_TEXT = [("신양리", "반가공센터"), ("반가공", "반가공센터"), ("가공 칠라", "반가공센터"), ("가공칠라", "반가공센터"),
                 ("액냉각", "접종·냉각실"), ("발이", "팽이 발이실"), ("생육", "팽이 생육실"), ("냉각", "접종·냉각실"), ("방냉", "접종·냉각실"),
                 ("접종", "접종·냉각실"), ("실험실", "실험실"), ("새송이배양", "새송이 배양실"), ("세송이배양", "새송이 배양실"),
                 ("배양", "팽이 배양실"), ("공조실", "팽이 배양실"), ("공조 ", "팽이 배양실"),
                 ("새송이", "새송이 재배사"), ("세송이", "새송이 재배사"), ("저장고", "포장·저장고"),
                 ("가마", "입병실"), ("살균", "입병실"), ("보일러", "입병실"), ("포장", "포장·저장고"),
                 ("탈병", "탈병실"), ("입병", "입병실"), ("긁기", "균긁기실"), ("권지", "권지실")]


def _area_from_text(text, default="냉동 공사(일백)"):
    """구역 시트가 아닌 줄(일백 시트)은 내용으로 구역을 짐작한다"""
    t = str(text).replace("미래공조", "")
    return next((a for k, a in _AREA_BY_TEXT if k in t), default)


_REPAIR_CAT = [                                    # 위에서부터 처음 맞는 분류 하나
    ("냉매·누수", ("냉매", "넹매", "냉메", "께스", "누수", "누설", "유불리기", "유블리기", "액불리기", "액블리기", "액분리기", "유분리기",
                  "자바라", "후렉시볼", "스트레이너", "스트레이나", "분배기", "비후다", "비우다", "비우라", "가스약간", "가스 조절", "가스조절")),
    ("팽창변·전자변", ("팽창변", "평창변", "펭창변", "전자변")),
    ("콤프", ("콤푸", "콤프")),
    ("패킹", ("페킹", "패킹")),
    ("오일", ("오일",)),
    ("실린더·센서", ("실린더", "센서", "쎈서", "썬서", "센거", "엔코더", "솔밸브", "쏠밸브", "에어쏠", "리데나", "오링", "로드셀", "감지기",
                    "압력스위치", "압력계")),
    ("전기·제어", ("기판", "PLC", "P.L.C", "마그네트", "메인보드", "단선", "전원", "계전기", "차단", "선절단", "신호선", "프로그램", "쎄팅", "세팅",
                  "인버터", "EOCR", "EMOR", "이오씨알", "OP걸림", "O.P", "op걸림", "lp트립", "LP트립", "프레샤", "플레샤", "M.C", "m.c", "단자함",
                  "리미트", "스위치", "등 교체", "등교체", "카메라", "터치")),
    ("필터", ("필터", "헤파", "FFU", "드라이어")),
    ("보일러·스팀", ("보일러", "보일", "스팀", "급수", "싸이폰")),
    ("물·배관", ("펌프", "지하수", "물탱크", "물라인", "수위", "배수", "감지봉")),
    ("도어", ("도어", "도아", "도오", "안쪽문", "바깥쪽문", "출입문", "이동문", "문 하한")),
    ("환기", ("덴파", "댐파", "단파", "시로코", "환기", "배기 펜", "배기펜", "배기 팬", "배기팬", "배기홴")),
    ("청소", ("청소",)),
    ("쿨러·팬·히터", ("쿨러", "쿨라", "홴", "팬", "펜", "에바", "히타", "희타", "히터")),
    ("차량·운반", ("지게차", "견인차", "렌탈", "고소작업", "바퀴", "베터리", "밧데리")),
    ("구동부", ("베어링", "베아링", "베오링", "베이링", "체인", "스프라켓", "모타", "모터", "머터", "감속기", "콘베아", "콘베어", "컨베아",
                "기어", "기아", "로라", "모차", "밸트", "벨트", "업소버", "텐션")),
    ("기계 부품", ("칼날", "타공", "환봉", "LM25", "LMK", "리본", "가이드", "집게", "볼트", "축", "파쇄기", "날개", "인펠라", "발판", "판",
                  "부러쉬", "봉")),
    ("건물 보수", ("에폭시", "몰탈", "방수", "천장")),
    ("설치·개선", ("설치", "변경", "보강", "추가", "철거", "조정", "조절")),
]
_REPAIR_EQUIP = [("저장고", ("저장고", "저온창고")), ("스피드 도어", ("스피드 도어", "스피드도어", "스피드 도아", "스피드도아")), ("권지세척기", ("권지세척기",)),
                 ("권지 건조기", ("권지 건조기", "권지건조기", "권지건저기")), ("입병기", ("입병기",)), ("하차기", ("하차기", "히차기")),
                 ("상차기", ("상차기",)), ("타공기", ("타공",)), ("접종기", ("접종기", "노즐")), ("탈병기", ("탈병기",)), ("파쇄기", ("파쇄기",)),
                 ("긁기기", ("긁기", "주수기")), ("콤프레샤", ("콤푸레샤", "콤프레샤")), ("크린부스", ("크린부스", "크린브스")),
                 ("지게차·견인차", ("지게차", "렌탈", "견인차", "고소작업차")), ("적재기", ("적재기", "적제기")), ("콘베아", ("콘베아",)),
                 ("에어컨", ("에어컨",)), ("가습기", ("가습기",)), ("실외기", ("실외기",)), ("냉동기", ("냉동기",))]


def _rep_nums(s):
    """'9,11,15' → [9, 11, 15] / '3~11' → 3..11"""
    out = []
    for part in re.split(r"[,.·]", str(s)):
        part = part.replace("번", "").strip()
        m = re.fullmatch(r"(\d{1,2})\s*[~]\s*(\d{1,2})", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a <= b <= a + 20:
                out += list(range(a, b + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def repair_places(text, area):
    """수리 내용에서 장소(여러 개)와 환경 기록 방 이름을 뽑는다. 반환 (장소 목록, 방 목록)"""
    t = str(text)
    t = re.sub(r"(?:\d\s*[,.]\s*)*\d\s*번\s*(?:살균\s*)?가마\s*\d\s*대로\s*가동", "", t)   # '1,3번 가마2대로 가동' 은 수리 장소가 아님
    places, rooms = [], []

    def put(p, r=()):
        if p not in places:
            places.append(p)
        for x in r:
            if x not in rooms:
                rooms.append(x)
    used = []                                       # 이미 쓴 글자 구간 (다른 규칙이 다시 읽지 않게)

    def spans(pat):
        for m in re.finditer(pat, t):
            if any(a <= m.start() < b for a, b in used):
                continue
            used.append(m.span())
            yield m
    NUMS = r"((?:\d{1,2}\s*번?\s*[,.]\s*)*\d{1,2})"   # '5번,7번' · '9,11,15' 같은 번호 목록
    ss_word = "새송이" in t or "세송이" in t
    # 냉각실·방냉실 (액냉각실·탱크 냉각실은 따로)
    for m in spans(r"액\s*냉각\s*실|(?:액\s*)?탱크\s*냉각\s*실"):
        put("액냉각실")
    for m in spans(r"(?:냉각\s*실?|방냉\s*실?)\s*(\d)\s*[-–]\s*(\d)"):
        r = "냉각1실" if m.group(1) == "1" else "냉각실2"
        put(r, [r])
    for m in spans(r"((?:\d\s*번?\s*[,.]\s*)*\d)\s*번?\s*냉각\s*실|(?:냉각\s*실|방냉\s*실)\s*((?:\d\s*번?\s*[,.]\s*)*\d)\s*번"):
        for n in _rep_nums(m.group(1) or m.group(2)):
            if n in (1, 2):
                r = "냉각1실" if n == 1 else "냉각실2"
                put(r, [r])
            else:
                put(f"냉각실 {n}번")
    if area == "접종·냉각실":                       # 접종실 시트의 '1-2번 냉동기' = 냉각1실 2번 냉동기
        for m in spans(r"(?<!\d)([12])\s*[-–]\s*\d(?!\d)"):
            r = "냉각1실" if m.group(1) == "1" else "냉각실2"
            put(r, [r])
    # 발이실 (범위 → 목록)
    for m in spans(r"발이\s*실?\s*(\d{1,2})\s*번?\s*~\s*(?:발이\s*실?\s*)?(\d{1,2})"):
        for n in range(int(m.group(1)), int(m.group(2)) + 1):
            if 1 <= n <= 18:
                put(f"발이실{n:02d}", [f"발이실{n:02d}"])
    for m in spans(r"발이\s*실?\s*" + NUMS):
        for n in _rep_nums(m.group(1)):
            if 1 <= n <= 18:
                put(f"발이실{n:02d}", [f"발이실{n:02d}"])
    # 생육실
    for m in spans(r"생육\s*실?\s*(\d{1,2})\s*번?\s*~\s*(?:생육\s*실?\s*)?(\d{1,2})"):
        for n in range(int(m.group(1)), int(m.group(2)) + 1):
            if 1 <= n <= 12:
                put(f"생육실{n:02d}", [f"생육실{n:02d}"])
    for m in spans(r"생육\s*실?\s*" + NUMS):
        for n in _rep_nums(m.group(1)):
            if 1 <= n <= 12:
                put(f"생육실{n:02d}", [f"생육실{n:02d}"])
    # 새송이 배양실 '1-3번 냉동기' · '3배양실3-1,2' (환경 기록 방이 없어 장소만)
    if area == "새송이 배양실" or re.search(r"(?:새송이|세송이)\s*배양", t):
        for m in spans(r"(?<!\d)(\d)\s*[-–]\s*((?:\d{1,2}\s*번?\s*[,.]\s*)*\d{1,2})(?!\d)"):
            for n in _rep_nums(m.group(2)):
                put(f"새송이 배양실 {m.group(1)}-{n}번")
    # 새송이 재배사 "1-3" = A03 (1=A, 2=B, 3=C, 4=D), "1동 4번" = A04, "4동" = D동
    ss_ctx = (ss_word or area == "새송이 재배사") and "배양" not in t
    for m in spans(r"(?:새송이|세송이)\s*(\d)\s*동\s*(\d{1,2})\s*번"):
        blk = C.SS_BLOCK.get(m.group(1))
        if blk:
            r = f"{blk}{int(m.group(2)):02d}"; put(r, [r])
    if ss_ctx:
        for m in spans(r"(?<!\d)(\d)\s*[-–]\s*(\d{1,2})(?!\d)"):
            blk = C.SS_BLOCK.get(m.group(1))
            if blk and 1 <= int(m.group(2)) <= 14:
                r = f"{blk}{int(m.group(2)):02d}"; put(r, [r])
        for m in spans(r"(?:새송이|세송이)\s*(\d)\s*동"):
            blk = C.SS_BLOCK.get(m.group(1))
            if blk:
                put(f"새송이 {blk}동")
    # 팽이 배양실: '1배양실'·'1,2번 배양실' = 방, '배양실 6번'·'6번 냉동기'·'공조실1번' = 냉동기 번호
    # (냉동기가 어느 방 것인지 기록이 없어, 방을 안 적었으면 배양실1·2 모두에 붙인다)
    if area not in ("새송이 배양실", "실험실") and not ss_word and "실험실" not in t and \
            (area == "팽이 배양실" or "배양" in t or re.search(r"공조\s*실?\s*\d", t)):
        bae = []
        for m in spans(r"((?:\d\s*번?\s*[,.]\s*)*\d)\s*번?\s*배양\s*실"):
            for n in _rep_nums(m.group(1)):
                if n in (1, 2):
                    put(f"배양실{n}", [f"배양실{n}"]); bae.append(f"배양실{n}")
        units = []
        for m in spans(r"(?:배양\s*실?|공조\s*실?)\s*" + NUMS + r"\s*(?:번|호기)"):
            units += _rep_nums(m.group(1))
        for m in spans(NUMS + r"\s*번\s*(?:냉동기|앤동기|유니트|쿨러|쿨라|콤푸|콤프|냉매|넹매|좌측|우측|실외기|오일|팽창변|평창변)"):
            units += _rep_nums(m.group(1))
        for n in dict.fromkeys(units):
            if 1 <= n <= 20:
                put(f"배양실 {n}번 냉동기", bae or ["배양실1", "배양실2"])
        if not places and (area == "팽이 배양실" or "배양" in t) and not re.search(r"냉각|방냉|예냉", t):
            put("팽이 배양실(공용)", ["배양실1", "배양실2"])
    # 살균솥(가마)
    for m in spans(r"((?:\d\s*[,.]\s*)*\d)\s*(?:번|호기?)\s*(?:살균\s*)?(?:가마|살균기|사균가마)"):
        for n in _rep_nums(m.group(1)):
            if 1 <= n <= 4:
                put(f"살균솥 {n}호기", [f"{n}호기"])
    for m in spans(r"(?:살균\s*가마|사균\s*가마|살균기|가마)\s*((?:\d\s*(?:번|호)?\s*)+)"):
        for n in re.findall(r"\d", m.group(1)):
            if 1 <= int(n) <= 4:
                put(f"살균솥 {int(n)}호기", [f"{int(n)}호기"])
    # 보일러
    for m in spans(r"(\d)\s*(?:번|호)\s*(?:가스\s*)?보일|보일러\s*(\d)\s*호기"):
        n = m.group(1) or m.group(2)
        put(f"보일러 {n}호기")
    # 배합기·혼합기
    for m in spans(r"(\d)\s*번?\s*(?:배합기|베합기|혼합기)"):
        put(f"배합기 {m.group(1)}번")
    for m in spans(r"(?:배합기|베합기|혼합기)\s*((?:\d\s*[,.]\s*)*\d)\s*호기"):
        for n in _rep_nums(m.group(1)):
            put(f"배합기 {n}번")
    # 저장고 냉동기
    for m in spans(r"(?:저장고|저온\s*창[고거])\s*(?:냉동기)?\s*(\d)\s*번"):
        put(("반가공 저온저장고" if "신양리" in t else f"저장고 {m.group(1)}번"))
    # 구역 시트 안에서 '6번방'·'11번 냉동기'처럼 구역 이름 없이 번호만 쓴 경우
    if not rooms and area in ("팽이 발이실", "팽이 생육실"):
        for m in spans(NUMS + r"\s*번\s*(?:방|냉동기|실외기|냉매|넹매|냉메|유불리기|유블리기|유압|홴|콤푸|콤프|쿨러|쿨라|팽창변|평창변|냉동|응축|고압)"):
            for n in _rep_nums(m.group(1)):
                lim = 18 if area == "팽이 발이실" else 12
                if 1 <= n <= lim:
                    r = f"{'발이실' if area == '팽이 발이실' else '생육실'}{n:02d}"; put(r, [r])
    # 냉각실 공용 (번호 없이 방 종류만)
    if not rooms and re.search(r"냉각|방냉|예냉", t) and not any("냉각" in p for p in places):
        put("냉각실(공용)", ["냉각1실", "냉각실2"])
    # 방·장비를 못 찾으면 장비 이름 → 구역 이름
    if not places:
        for name, keys in _REPAIR_EQUIP:
            if any(k in t for k in keys):
                put(f"{area} · {name}")
                break
    if not places:
        put(area)
    return places, rooms


_REFRIG = r"냉동기|앤동기|냉각|방냉|발이|생육|배양|공조|쿨러|쿨라|냉매|넹매|실외기|유니트|에바|응축"


def repair_category(text, area=""):
    t = str(text)
    if re.search(r"입고\W*$", t.strip()):
        return "부품·장비 입고"                         # 냉매 5통 입고 · 홴모타 입고 같은 줄
    if area == "콤프레샤" or (re.search(r"에어\s*콤|에오\s*콤|콤[푸프]레샤", t) and not re.search(_REFRIG, t)):
        return "에어 콤프레샤"                          # 공장 압축공기용 (냉동기 콤프와 구분)
    for cat, keys in _REPAIR_CAT:
        if cat == "쿨러·팬·히터" and "용접" in t and re.search(_REFRIG + r"|배관", t):
            return "냉매·누수"                          # 쿨러·배관 용접 = 냉매가 새는 곳을 막는 일
        if any(k in t for k in keys):
            return cat
    if "용접" in t and re.search(_REFRIG + r"|배관", t):
        return "냉매·누수"
    if "정검" in t or "점검" in t:
        return "점검"
    return "기타"


def _repair_vendor(*texts):
    s = " ".join(str(x) for x in texts if x)
    for v, keys in (("일백", ("일백", "앨백", "노기사", "이기사")), ("디엠에스", ("디엠에스", "DMS", "dms", "D.M.S")), ("세양", ("세양",)),
                    ("미래공조", ("미래공조",)), ("퓨어시스템", ("퓨어시스템",))):
        if any(k in s for k in keys):
            return v
    return ""


def _repair_date(v, today):
    if isinstance(v, dt.datetime):
        d = v.date()
    elif isinstance(v, dt.date):
        d = v
    elif isinstance(v, str):
        m = re.match(r"\s*(20\d{2})[-./]\s*(\d{1,2})[-./]\s*(\d{1,2})", v)
        if not m:
            return None, False
        try:
            d = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None, False
    else:
        return None, False
    if d.year < 2015:
        return None, False
    fixed = False
    if d > today + dt.timedelta(days=7):                 # 연도를 잘못 쓴 줄(예: 2026-12-19 → 2025-12-19)
        try:
            d2 = d.replace(year=d.year - 1)
            if d2 <= today:
                d, fixed = d2, True
        except ValueError:
            pass
    return d, fixed


def _rep_norm(s):
    return re.sub(r"[^0-9A-Za-z가-힣]", "", str(s))


def _rep_specific(places, area):
    """방·호기 수준으로 콕 집은 장소만 (구역 이름, '구역 · 장비', '(공용)', '새송이 A동' 은 빼고)"""
    return {p for p in places if p != area and " · " not in p and "(공용)" not in p and not re.fullmatch(r"새송이 [A-D]동", p)}


def _rep_same(k, rec, difflib):
    """같은 날 두 줄이 같은 수리인지"""
    if k["_n"] == rec["_n"]:
        return True
    if k["시트"] == rec["시트"]:
        return False
    sa, sb = k["_spec"], rec["_spec"]
    if sa and sb and not (sa & sb):
        return False
    ratio = difflib.SequenceMatcher(None, k["_n"], rec["_n"]).ratio()
    if {k["시트"], rec["시트"]} & {"일백", "기타장비"}:
        if ratio >= 0.72 or (bool(sa & sb) and ratio >= 0.5):
            return True
        # 한쪽이 방을 못 찾은 줄이면 '6번 냉동기'·'2-2' 같은 번호가 같고 절반 이상 닮았을 때
        return (not sa or not sb) and bool(_rep_numtok(k["내용"]) & _rep_numtok(rec["내용"])) and ratio >= 0.5
    return ratio >= 0.9 and sa == sb


def _rep_numtok(text):
    t = str(text)
    return set(re.findall(r"(\d{1,2})\s*번", t)) | {f"{a}-{b}" for a, b in re.findall(r"(?<!\d)(\d)\s*[-–]\s*(\d{1,2})(?!\d)", t)}


def parse_repairs(path, today=None):
    """기계수리.xlsx → {'repairs': DataFrame[date, 구역, 장소, 방, 분류, 내용, 사양, 업체, 비고, 시트, 행]}"""
    import difflib
    today = today or dt.date.today()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rank = {"기타장비": 1, "일백": 2}                    # 구역 시트를 먼저 읽고, 모아 적은 시트(기타장비·일백)는 나중에 → 겹치면 구역 시트 줄을 남긴다
    names = sorted(wb.sheetnames, key=lambda n: rank.get(n.strip(), 0))
    rows = []
    for sn in names:
        ws = wb[sn]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = dcol = tcol = scol = None
        for i, r in enumerate(grid[:10]):
            txt = [str(v).strip() if v is not None else "" for v in r]
            if any(x in ("날짜", "날자") for x in txt) and any("내역" in x for x in txt):
                hdr = i; dcol = next(j for j, x in enumerate(txt) if x in ("날짜", "날자"))
                tcol = next(j for j, x in enumerate(txt) if "내역" in x)
                scol = next((j for j, x in enumerate(txt) if x == "사양"), None)
                break
        if hdr is None:
            continue
        area0 = REPAIR_AREA.get(sn.strip(), sn.strip())
        last = None
        tend = scol if scol is not None else tcol + 4
        for i in range(hdr + 1, len(grid)):
            r = grid[i] + [None] * 3
            d, fixed = _repair_date(r[dcol], today)
            text = " ".join(str(v).strip() for v in r[tcol:tend] if isinstance(v, str) and v.strip())
            spec = str(r[scol]).strip() if scol is not None and r[scol] not in (None, "") else ""
            if d is None and r[dcol] not in (None, "") and not isinstance(r[dcol], str):
                d = None                                  # 숫자만 있는 칸(번호) 등
            if d is not None:
                last = d
            if text and (d or last):
                use_d = d or last
                area = area0 or _area_from_text(text)
                rows.append({"date": use_d, "구역": area, "내용": re.sub(r"\s+", " ", text), "사양": spec,
                             "업체": _repair_vendor(text, spec) or ("일백" if sn.strip() == "일백" else ""),
                             "비고": "날짜 연도 보정" if fixed else "", "시트": sn.strip(), "행": i + 1})
            # 같은 줄 오른쪽에 따로 적은 '날짜 | 내용' 쌍 (예: 일백 시트의 입고 메모)
            for j in range((scol if scol is not None else tend) + 1, len(r) - 1):
                d2, _ = _repair_date(r[j], today)
                if d2 and isinstance(r[j + 1], str) and r[j + 1].strip():
                    t2 = r[j + 1].strip()
                    area = area0 or _area_from_text(t2)
                    rows.append({"date": d2, "구역": area, "내용": re.sub(r"\s+", " ", t2), "사양": "", "업체": _repair_vendor(t2) or ("일백" if sn.strip() == "일백" else ""),
                                 "비고": "", "시트": sn.strip(), "행": i + 1})
    wb.close()
    # 같은 일을 두 시트에 적은 줄(구역 시트 ↔ 일백·기타장비 시트)은 한 줄로 합친다.
    #  - 같은 시트 안: 글자가 똑같을 때만 (6번·5번 냉동기, 9·10·11번 실외기처럼 비슷해도 다른 일)
    #  - 서로 다른 방·장비를 가리키면 합치지 않는다
    #  - 일백·기타장비 시트 ↔ 구역 시트: 내용이 많이 닮았거나(72% 이상), 같은 방이고 절반 이상 닮았으면 합친다
    #  - 구역 시트끼리: 거의 같은 글(90% 이상)이고 가리키는 방·장비가 같을 때만
    keep = []
    by_date = {}
    for rec in rows:
        places, rms = repair_places(rec["내용"], rec["구역"])
        rec["_places"], rec["_rooms"] = places, rms
        rec["_spec"] = _rep_specific(places, rec["구역"])
        rec["_n"] = _rep_norm(rec["내용"])
        dup = None
        for k in by_date.get(rec["date"], []):
            if _rep_same(k, rec, difflib):
                dup = k; break
        if dup is not None:
            if rec["업체"] and not dup["업체"]:
                dup["업체"] = rec["업체"]
            if rec["사양"] and rec["사양"] not in dup["사양"]:
                dup["사양"] = (dup["사양"] + " · " if dup["사양"] else "") + rec["사양"]
            if rec["_n"] != dup["_n"]:
                note = f"{rec['시트']} 시트: {rec['내용']}"
                dup["비고"] = (dup["비고"] + " / " if dup["비고"] else "") + note
            if rec["_spec"] and not dup["_spec"]:          # 남긴 줄은 방을 안 적었고 합친 줄이 적었으면 그 방을 쓴다
                dup["_places"] = [pl for pl in rec["_places"] if pl in rec["_spec"]]
                dup["_rooms"] = list(rec["_rooms"]); dup["_spec"] = set(rec["_spec"])
            elif rec["_spec"]:                               # 둘 다 적었으면 빠진 방만 보탠다
                for pl in rec["_places"]:
                    if pl in rec["_spec"] and pl not in dup["_places"]:
                        dup["_places"].append(pl); dup["_spec"].add(pl)
                for rm in rec["_rooms"]:
                    if rm not in dup["_rooms"]:
                        dup["_rooms"].append(rm)
            continue
        by_date.setdefault(rec["date"], []).append(rec)
        keep.append(rec)
    out = []
    for rec in keep:
        out.append({"date": pd.Timestamp(rec["date"]), "구역": rec["구역"], "장소": ", ".join(rec["_places"]), "방": ", ".join(rec["_rooms"]),
                    "분류": repair_category(rec["내용"], rec["구역"]), "내용": rec["내용"], "사양": rec["사양"], "업체": rec["업체"],
                    "비고": rec["비고"], "시트": rec["시트"], "행": int(rec["행"])})
    df = pd.DataFrame(out)
    if not df.empty:
        df = df.sort_values(["date", "구역", "행"]).reset_index(drop=True)
    return {"repairs": df}


def parse_humidifier_repairs(path, today=None):
    """가습기 수리내역.xlsx (월별 시트, 날짜 × 부품별 교체 수) → repairs 행 (방 기록이 없어 장소는 '팽이 가습기')"""
    today = today or dt.date.today()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for sn in wb.sheetnames:
        grid = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        hdr = None
        for i, r in enumerate(grid[:8]):
            txt = [str(v).strip() if v is not None else "" for v in r]
            if any("진동자" in x for x in txt):
                hdr = i; cols = {j: x for j, x in enumerate(txt) if x}
                break
        if hdr is None:
            continue
        room_col = next((j for j, x in cols.items() if "재배실" in x), None)
        part_cols = {j: x for j, x in cols.items() if j != room_col and j > 0}
        width = max(len(r) for r in grid)
        for i in range(hdr + 1, len(grid)):
            r = grid[i] + [None] * (width - len(grid[i]))
            d, _ = _repair_date(r[0], today)
            if d is None:
                continue
            parts = []
            for j in range(1, width):
                if j == room_col or r[j] in (None, ""):
                    continue
                name = part_cols.get(j, "기타")
                v = r[j]
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    if v:
                        parts.append(f"{name.replace(' ', '')} {int(v) if float(v).is_integer() else v}")
                elif isinstance(v, str) and v.strip():
                    parts.append(v.strip() if name == "기타" else f"{name.replace(' ', '')} {v.strip()}")
            if not parts:
                continue
            room = r[room_col] if room_col is not None else None
            ids = re.findall(r"(?:YJ)?\d{5}", str(room or "").replace(" ", ""))      # 'YJ25164YJ25166' · '2501825012' 처럼 붙여 쓴 번호도 나눈다
            note = ("가습기 번호: " + ", ".join(ids)) if ids else (f"재배실 칸: {room}" if room not in (None, "") else "")
            out.append({"date": pd.Timestamp(d), "구역": "팽이 재배사(가습기)", "장소": "팽이 가습기", "방": "",
                        "분류": "가습기 부품", "내용": "가습기 부품 교체 — " + ", ".join(parts), "사양": "",
                        "업체": "", "비고": note, "시트": sn.strip(), "행": i + 1})
    wb.close()
    df = pd.DataFrame(out)
    if not df.empty:
        df = df.sort_values("date").reset_index(drop=True)
    return {"repairs": df}


def _xlrd():
    """옛 엑셀(.xls)을 읽는 xlrd. 없으면 설치를 시도한다 — 하루에 한 번만(갱신 때마다 기다리지 않게)."""
    try:
        import xlrd
        return xlrd
    except ImportError:
        pass
    mark = os.path.join(C.CACHE_DIR, "_xlrd_install_try.txt")
    try:
        if os.path.exists(mark) and dt.datetime.now().timestamp() - os.path.getmtime(mark) < 86400:
            return None
        with open(mark, "w", encoding="utf-8") as fh:
            fh.write(dt.datetime.now().isoformat())
    except OSError:
        pass
    import subprocess, sys, importlib
    try:
        subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "xlrd"], timeout=120, check=False)
        importlib.invalidate_caches()
        import xlrd
        return xlrd
    except Exception:
        return None


def _sheets_any(path):
    """{시트 이름: 2차원 목록} — .xls 는 xlrd, .xlsx 는 openpyxl (날짜 칸은 datetime 으로)"""
    if str(path).lower().endswith(".xls"):
        xlrd = _xlrd()
        if xlrd is None:
            raise ImportError("xlrd 패키지가 없어 .xls 파일을 읽지 못했습니다 (pip install xlrd)")
        wb = xlrd.open_workbook(path)
        out = {}
        for sh in wb.sheets():
            g = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate_as_datetime(v, wb.datemode)
                        except Exception:
                            pass
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        v = None
                    row.append(v)
                g.append(row)
            out[sh.name] = g
        return out
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {sn: [list(r) for r in wb[sn].iter_rows(values_only=True)] for sn in wb.sheetnames}
    wb.close()
    return out


def _won(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    m = re.search(r"\d[\d,]*", str(v or ""))
    if m:
        try:
            return float(m.group(0).replace(",", ""))
        except ValueError:
            return None
    return None


def parse_repair_reports(path, today=None):
    """수리 완료 보고서(.xls) → {'repair_costs': DataFrame[date, 완료일, 구역, 수리명, 업체, 수리비, 발주처, 인원, 기간, 내용, 시트, 행, 열]}
    한 시트에 보고서 양식이 가로·세로로 여러 개 있다 → '수 리 명' 칸마다 한 건."""
    today = today or dt.date.today()
    out = []
    for sn, g in _sheets_any(path).items():
        H = len(g)

        def at(rr, cc):
            return g[rr][cc] if 0 <= rr < H and 0 <= cc < len(g[rr]) else None

        def lab(v):
            return re.sub(r"\s", "", v) if isinstance(v, str) else ""
        for r in range(H):
            for c, v in enumerate(g[r]):
                if lab(v) != "수리명":
                    continue

                def right_of(rr, label):
                    for cc in range(c, c + 8):
                        if lab(at(rr, cc)).startswith(label):
                            for c2 in range(cc + 1, cc + 4):
                                y = at(rr, c2)
                                if y not in (None, ""):
                                    return y
                    return None
                name = right_of(r, "수리명"); order = right_of(r, "발주처")
                cost = _won(right_of(r + 1, "수리비")); vendor = right_of(r + 1, "수리업체")
                d_rep, _ = _repair_date(right_of(r + 2, "수리일"), today)
                d_end, _ = _repair_date(right_of(r + 2, "수리완료일"), today)
                people = period = None
                items, in_txt, blank = [], False, 0
                for rr in range(r + 3, min(r + 32, H)):
                    l0 = lab(at(rr, c))
                    if l0.startswith("현장감독") or l0 == "수리명":
                        break
                    if "공사인원" in l0:
                        people = at(rr, c + 1)
                    elif "공사기간" in l0:
                        period = at(rr, c + 1)
                    elif "수리내용" in l0:
                        in_txt = True
                    if not in_txt:
                        continue
                    x = at(rr, c + 1)
                    x = re.sub(r"\s+", " ", x).strip() if isinstance(x, str) else ""
                    if not x:
                        blank += 1
                        if blank >= 3:
                            break
                        continue
                    blank = 0
                    if "견적서" in x:
                        break
                    if items and not re.match(r"^\d", x) and re.search(r"(?:으로|하여|및|,|고|후|는|을|를|의|에)$", items[-1]):
                        items[-1] += " " + x                     # 앞 줄이 말을 끝맺지 않았으면 이어 붙인다
                    else:
                        items.append(x)
                d = d_rep or d_end
                if d is None and not items:
                    continue
                text = " / ".join(items)
                vend = str(vendor or "").strip()
                area = "냉동 공사(일백)" if "일백" in vend else (_area_from_text(str(name or ""), default=None)
                                                                 or _area_from_text(text, default="기타"))
                out.append({"date": pd.Timestamp(d) if d else pd.NaT, "완료일": pd.Timestamp(d_end) if d_end else pd.NaT,
                            "구역": area, "수리명": str(name or "").strip(), "업체": vend, "수리비": cost,
                            "발주처": str(order or "").strip(), "인원": str(people or "").strip(), "기간": str(period or "").strip(),
                            "내용": text, "시트": str(sn).strip(), "행": r + 1, "열": c + 1})
    df = pd.DataFrame(out)
    if not df.empty:
        df = df.sort_values(["date", "시트", "행", "열"]).reset_index(drop=True)
    return {"repair_costs": df}


# ───────────────────────── 실행 ─────────────────────────
# (종류, 파일 패턴, 파서, 하위폴더 검색 여부)
SOURCES = [
    ("paengi", "*팽이 생산 상세일지.xlsx", parse_paengi_detail, True),
    ("paengi", "*팽이 일일작업일지.xlsx", parse_paengi_work, True),
    ("ss", "*새송이 일일작업일지.xlsx", parse_ss_work, True),
    ("ss", "*새송이 상세일지.xlsx", parse_ss_detail, True),
    ("ss_extra", "새송이 방별 포장량.xlsx", parse_ss_room_pack, False),
    ("cult", "*배양소관리.xlsx", parse_cult, True),
    # 원가 자료 (재배 관리자 폴더)
    ("kwon_paengi", "*팽이재배일보.xlsx", parse_bottling, False),
    ("kwon_ss", "*새송이작업일보.xlsx", parse_bottling, False),
    ("kwon_ss", "원재료입고.xlsx", parse_materials, False),
    ("kwon_paengi", "전기세.xlsx", parse_electric, False),
    ("kwon_paengi", "*가스 및 경유사용량*.xlsx", parse_fuel, False),
    # 종균·액탱크 (배양센터 최상위)
    ("cult_extra", "종균 관리 대장*.xlsx", parse_spawn_ledger, False),
    ("cult_extra", "배양실 작업 일지*.xlsx", parse_cult_worklog, False),
    # 재배 관리자 통합 일지 (솥 단위 입병 기록·매출)
    ("jung_paengi", "*팽이버섯 재배관리.xlsx", parse_master, False),
    ("jung_ss", "*새송이 일지.xlsx", parse_master, False),
    ("jung_ss", "생산량*.xlsx", parse_ss_room_qc, False),
    # 수리 기록 (권석현 폴더 — 새송이·팽이 어느 쪽에 두어도 읽는다. 같은 내용이 두 곳에 있으면 한 번만 남긴다)
    ("kwon_ss", "기계수리*.xlsx", parse_repairs, False),
    ("kwon_paengi", "기계수리*.xlsx", parse_repairs, False),
    ("kwon_ss", "가습기 수리*.xlsx", parse_humidifier_repairs, False),
    ("kwon_paengi", "가습기 수리*.xlsx", parse_humidifier_repairs, False),
    ("kwon_ss", "수리 완료 보고서*.xls*", parse_repair_reports, False),
    ("kwon_paengi", "수리 완료 보고서*.xls*", parse_repair_reports, False),
]


def _too_old(path):
    """경로에 'YYYY년' 폴더가 있으면 LOG_MIN_YEAR 미만은 건너뜀."""
    for part in path.replace("\\", "/").split("/"):
        m = re.fullmatch(r"(\d{4})년", part.strip())
        if m and int(m.group(1)) < C.LOG_MIN_YEAR:
            return True
    return False


def source_files():
    """(종류, 파일경로, 파서) 목록. 원본 폴더는 config.LOG_SOURCES 순서로 찾는다."""
    out = []
    for kind, pattern, fn, rec in SOURCES:
        base = C.log_dir(kind)
        pat = os.path.join(base, "**", pattern) if rec else os.path.join(base, pattern)
        for path in sorted(glob.glob(pat, recursive=rec)):
            name = os.path.basename(path)
            rel = os.path.relpath(path, base)
            if name.startswith("~$") or _too_old(rel) or any(k in rel for k in C.LOG_SKIP_DIRS):
                continue
            out.append((kind, path, fn))
    return out


def run(progress=None, force=False):
    C.ensure_dirs()
    st = json.load(open(STATE, encoding="utf-8")) if os.path.exists(STATE) and not force else {}
    data = {k: (pd.read_parquet(_p(k)) if os.path.exists(_p(k)) and not force else pd.DataFrame()) for k in DATASETS}
    changed, n = set(), 0
    live = set()
    for kind, path, fn in source_files():
        if True:
            base = os.path.basename(path)
            key = path
            live.add(key)
            mt = os.path.getmtime(path)
            if st.get(key, {}).get("mtime") == mt:
                continue
            try:
                res = fn(path)
            except Exception as e:                       # 한 파일이 깨져도 나머지는 진행
                # 패키지가 없어 못 읽은 파일(.xls 의 xlrd 등)은 수정시각을 남기지 않아 다음 갱신 때 다시 시도한다
                st[key] = {"mtime": None if isinstance(e, ImportError) else mt, "error": f"{type(e).__name__}: {e}"}
                if progress:
                    progress(f"오류 {base}: {e}")
                continue
            for k, df in res.items():
                if df is None or df.empty:
                    continue
                df = df.copy(); df["_src"] = key
                old = data[k]
                if not old.empty and "_src" in old:
                    old = old[old._src != key]
                data[k] = pd.concat([old, df], ignore_index=True)
                changed.add(k)
            st[key] = {"mtime": mt, "rows": {k: int(len(v)) for k, v in res.items()}, "read": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M")}
            n += 1
            if progress:
                progress(f"읽음 {base}")
    gone = [k for k in st if k not in live and not k.startswith("_")]
    for k in gone:
        for ds in DATASETS:
            if not data[ds].empty and "_src" in data[ds] and (data[ds]._src == k).any():
                data[ds] = data[ds][data[ds]._src != k]; changed.add(ds)
        st.pop(k, None)
    for k in changed:
        df = data[k]
        # 같은 날짜가 두 파일에 있으면 나중에 읽은 것을 남긴다
        dedup = {"paengi_daily": ["date"], "paengi_harvest": ["date"], "paengi_qc": ["date", "room", "kind"],
                 "ss_daily": ["date"], "ss_events": ["date", "event", "room"], "ss_kgcat": ["date"],
                 "ss_room_pack": ["date", "room"], "ss_room_in": ["입상일", "room"], "ss_room_qc": ["입상일", "room"],
                 "bottling": ["date", "crop"], "materials": ["month", "item"], "materials_block": ["month"],
                 "electric": ["month", "meter"], "fuel": ["date", "fuel"], "ledger": ["crop", "month", "seq"],
                 "cult_inoc": ["일자", "품목"], "cult_scrape_paengi": ["긁기일", "방", "접종일"],
                 "cult_scrape_ss": ["긁기일", "방", "접종일"], "cult_kwonji": ["권지일", "생육실", "발이실"],
                 "spawn_tanks": ["코드", "탱크NO"], "tank_daily": ["date"], "loads": ["crop", "입병일", "입병순번"], "sales": ["crop", "month"],
                 "repairs": ["date", "시트", "내용"], "repair_costs": ["date", "시트", "수리명", "수리비", "내용"]}[k]
        df = df.drop_duplicates(dedup, keep="last")
        for c in df.columns:
            if df[c].dtype == object:
                df[c] = df[c].astype(str).where(df[c].notna(), None)
        df.to_parquet(_p(k) + ".tmp", index=False); os.replace(_p(k) + ".tmp", _p(k))
    json.dump(st, open(STATE, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    return n


def load(name):
    p = _p(name)
    if os.path.exists(p):
        df = pd.read_parquet(p)
        for c in df.columns:
            if c in ("date", "긁기일", "권지일", "일자", "접종일", "입상일", "일시", "입병일", "균긁기일", "수확일", "솎기일", "완료일"):
                df[c] = pd.to_datetime(df[c], errors="coerce")
        return df
    return pd.DataFrame()


def status():
    return json.load(open(STATE, encoding="utf-8")) if os.path.exists(STATE) else {}
