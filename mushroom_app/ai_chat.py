# -*- coding: utf-8 -*-
"""AI 질문창: Claude API 로 앱의 데이터(캐시 표)와 일지 파일을 찾아 답한다.
- API 키: 환경변수 ANTHROPIC_API_KEY 또는 DATA_DIR\\anthropic_key.txt (첫 줄)
- 도구: 표 설명 / pandas 코드 실행(읽기 전용) / 파일 찾기 / 엑셀 미리보기
"""
import os, io, re, glob, json, contextlib, traceback
import numpy as np
import pandas as pd
from . import config as C

KEY_FILE = os.path.join(C.DATA_DIR, "anthropic_key.txt")
MODEL = getattr(C, "AI_MODEL", "claude-sonnet-4-5")
MAX_TURNS = 8                     # 한 질문당 도구 호출 최대 횟수
MAX_OUT = 6000                    # 도구 결과 최대 글자수

# 파일 도구가 볼 수 있는 폴더 (앱 데이터 폴더 아래 + Z: 원본 일지 폴더)
def _roots():
    r = [C.DATA_DIR]
    for k, lst in C.LOG_SOURCES.items():
        r.append(lst[0])
    return [os.path.abspath(p) for p in r if os.path.isdir(p)]


def _key_candidates():
    """키 파일을 찾아볼 경로 목록 (순서대로). 파일 내용은 찾을 때 읽지 않는다."""
    names = getattr(C, "AI_KEY_NAMES", ["anthropic_key.txt"])
    out = []
    exact = (getattr(C, "AI_KEY_FILE", "") or "").strip()
    if exact:
        out.append(exact)
    for n in names:
        out.append(os.path.join(C.DATA_DIR, n))
    # 앱 폴더 (…\unifarm\app\) — 이 파일이 있는 mushroom_app 의 한 단계 위
    appdir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out += [os.path.join(appdir, n) for n in names]
    # Downloads 바로 아래 폴더 한 겹까지 (예: Downloads\astra_app\claude_api_key.txt)
    dl = os.path.dirname(os.path.abspath(C.DATA_DIR))
    try:
        for d in sorted(os.listdir(dl)):
            sub = os.path.join(dl, d)
            if os.path.isdir(sub):
                out += [os.path.join(sub, n) for n in names]
        out += [os.path.join(dl, n) for n in names]
    except OSError:
        pass
    return out


def key_path():
    """실제로 쓰는 키 파일 경로. 없으면 None. (경로만 돌려주고 내용은 담지 않는다)"""
    for p in _key_candidates():
        if p and os.path.isfile(p):
            return p
    return None


def api_key():
    k = os.environ.get("ANTHROPIC_API_KEY", "").strip()
    if k:
        return k
    p = key_path()
    if p:
        with open(p, encoding="utf-8") as f:
            k = f.readline().strip()
    return k or None


# ───────────────────────── 도구 ─────────────────────────
TOOLS = [
    {"name": "describe_tables",
     "description": "앱이 들고 있는 데이터 표들의 이름·열·행수·기간을 알려준다. 데이터 질문이면 먼저 이걸 불러 어떤 표를 쓸지 정한다.",
     "input_schema": {"type": "object", "properties": {}}},
    {"name": "run_pandas",
     "description": ("pandas 코드를 실행해 결과를 돌려준다. 표는 딕셔너리 T 로 접근한다 (예: T['paengi_batches']). pd, np 사용 가능. "
                     "마지막 줄의 식(expression) 값 또는 print() 출력이 결과로 돌아온다. 파일 쓰기·네트워크·import 는 금지."),
     "input_schema": {"type": "object", "properties": {"code": {"type": "string", "description": "실행할 파이썬 코드"}}, "required": ["code"]}},
    {"name": "find_files",
     "description": "일지·자료 폴더에서 파일 이름에 키워드가 들어간 파일을 찾는다 (엑셀·CSV·PDF 등). 여러 키워드는 공백으로 구분(모두 포함).",
     "input_schema": {"type": "object", "properties": {"keywords": {"type": "string"}, "limit": {"type": "integer", "default": 30}}, "required": ["keywords"]}},
    {"name": "read_excel",
     "description": "엑셀 파일의 시트 목록 또는 특정 시트의 앞부분(최대 60행)을 읽는다. sheet 를 비우면 시트 목록과 첫 시트 미리보기.",
     "input_schema": {"type": "object", "properties": {"path": {"type": "string"}, "sheet": {"type": "string"}, "max_rows": {"type": "integer", "default": 40},
                                                       "skip_rows": {"type": "integer", "default": 0}}, "required": ["path"]}},
]

_FORBID = re.compile(r"(__|import\s|open\s*\(|exec\s*\(|eval\s*\(|os\.|sys\.|subprocess|shutil|to_csv|to_excel|to_parquet|\.write\(|pathlib|socket|requests)")


def _tables(loaders):
    """loaders: {이름: 함수} → 지연 로딩된 표 딕셔너리."""
    class Lazy:
        def __init__(self):
            self._c = {}
        def __getitem__(self, k):
            if k not in self._c:
                if k not in loaders:
                    raise KeyError(f"표 없음: {k}. describe_tables 로 목록을 확인하세요.")
                self._c[k] = loaders[k]()
            return self._c[k]
        def keys(self):
            return list(loaders.keys())
        def __contains__(self, k):
            return k in loaders
    return Lazy()


TABLE_DOC = {
    "env": "재배동 방별·날짜별 환경 일집계 (farm, room, date, crop, run 운전비율, temp/hum/co2 실측 평균, sp_temp/sp_hum/sp_co2 설정 평균, cool/heat/humid 가동비율, track_err 추종오차, day_last 재배일차 등)",
    "paengi_batches": "팽이 배치 사슬 (bid, 접종일, 입상일=균긁기, 발이실, 권지일, 생육실, 수확시작/종료, 총일수, 배양일수, 병당중량_표시, 중량기준, QC, 긁기불량률%, 권지불량률%, 접종손실률%, 탱크, season, month, 완결)",
    "paengi_batch_env": "팽이 배치별 일차 환경 (bid, phase 배양/발이실/생육실, room, date, cday 일차, temp, sp_temp, hum, co2, cool, heat, humid, run)",
    "ss_batches": "새송이 배치 (bid, room, block, 입상일, bottles 입상병수, 뒤집기_d, 솎기_d, 수확시작_d, 수확일수, 재입상_d, 접종일, 배양일수, 수확kg, 병당g, 미니비중% (2025-10~2026-04 방별 포장량 자료 있는 배치만), season, month, 완결)",
    "ss_batch_env": "새송이 배치별 일차 환경 (bid, phase, room, date, cday, temp, sp_temp, hum, co2, cool, heat, humid, run)",
    "paengi_monthly": "팽이 월별 (월, 생산kg, 입상병수, 생산일, 병당중량, 상등급비중, QC)",
    "ss_monthly": "새송이 월별 (월, 생산kg, 입상병수, 배치병수, 병당g, 파지%, 미니%, 깎기%, 수확소요일, 재입상간격)",
    "cult_monthly": "배양실 월별 손실·불량률 (m, 항목, 비율%, 분자, 분모)",
    "cult_scrape_paengi": "팽이 균긁기 검수 (긁기일, 접종일, 배양출고, 불량, 방=발이실, 탱크)",
    "cult_scrape_ss": "새송이 균긁기 검수 (긁기일, 접종일, 배양출고, 불량, 방_raw)",
    "cult_kwonji": "팽이 권지 작업 (권지일, 접종일, 입상일, 발이실, 생육실, 입상량, 불량)",
    "cult_inoc": "접종 (일자, 품목, 접종량, 손실)",
    "paengi_harvest": "팽이 일별 수확 일지 (date, harvest_rooms 수확 생육실, bottles, g_per_bottle 병당중량, grade_top)",
    "paengi_qc": "팽이 QC 점수 (date, room 생육실, day_n, score, kind 본수확/점검수확, note, fmt 서식A/B)",
    "ss_daily": "새송이 일별 (date, kg 생산량, bottles 입상병수, 입상/수확/폐상 방 표기)",
    "ss_room_pack": "새송이 방별 일별 포장량 (date, room, 미니kg, 새송이kg, kg, 박스수) 2025-10~2026-04",
    "anomalies": "최근 60일 이상 감시 결과 (심각도, 분류, 재배사, 방, 날짜, 내용, 값)",
    "tank_rank": "팽이 액탱크별 균긁기 불량률 순위 (index 탱크, 배치, 불량률%, 단독배치, 단독불량률%, 최근90일%, 월평균대비pp)",
    "cost_monthly": "월·작물별 병당 원가 (월, 작물, 입병수, 입병일수, 원재료비, 연료비, 배양전기, 생육전기, 합계 [원/병], *_원 은 월 금액, 보고서_합계 는 월말보고 참고값)",
    "materials": "원재료 월별 품목별 (month, item, 전월재고, 입고량, 사용량, 현재고, 단가 원/kg, 팽이_kg, 새송이_kg, 팽이_원, 새송이_원)",
    "electric": "전기요금 월·계량기별 (month, meter, kwh, cost 원)",
    "fuel": "LPG·경유 충진 (date, fuel, qty kg/L, unit_price, amount 원)",
    "bottling": "재배일보 입병 (date, crop, bottles 입병수, sterilized 살균수, batches 회)",
    "ledger": "생산 상세일지 입출고 원장 (crop, month, date, type 재고/입고/출고, cust 거래처, sub 거래처세부, grade, name 규격명(새송이), spec 규격, box, kg, price 단가, amt 금액, note)",
    "stock_recon_paengi": "팽이 월별 재고 대사 (월, 기초재고, 입고, 출고, 계산기말, 일평균출고, 기말재고일수, 다음달기초, 월초조정, 조정률%)",
    "stock_recon_ss": "새송이 월별 재고 대사 (열은 팽이와 같음)",
    "spawn_tanks": "종균 관리 대장의 액탱크 배치 (일시 제조일, 품종, 탱크NO, 코드, 배치, 계통, 세대, 모균, 유래 조직분리 정보, 접종일, 교반, pH, 액배양일)",
    "tank_daily": "배양실 작업 일지 일별 (date, 접종탱크NO, 종균pH, 접종탱크코드, 접종배지, 살균후pH, 입병, 살균전pH, 삼각접종원, 신규탱크, 비고 설비·특이사항 메모)",
    "loads": "재배 관리자 종합관리 — 살균솥 한 번(혼합기 배치, 약 8,600병) 단위 기록 (crop, 입병일, 입병순번, 혼합기, 살균기NO, 입병시작/종료 분, 살균수량, pH, 수분, 중량상자, 충진시간, 배지품온, 함량, 접종일, 종균, 접종온도, 접종LOSS, 균긁기일, 배양일수, 발이실, 긁기불량, 입상수량, 긁기불량률%(균긁기 기록 행 단위 = 긁기불량÷(긁기불량+입상수량); 균긁기·권지 칸은 솥별이 아니라 그날 병 중 그 발이실로 간 몫이라 불량률은 입병일(하루) 단위로 볼 것), 긁기비고, 팽이: 권지일/생육실/권지불량/권지수량/권지불량률%(권지불량÷(권지불량+권지수량))/상태/수확일/평균중량, 새송이: 솎기일/수확일/생육기간; 살균 연결 시 살균시작/F0/2차유지분/중단횟수)",
    "sales": "월별 매출 (작물, 월, 출고kg, 매출원, 단가kg, 단가박스, 수확병수, 병당g, 입병수, 병당매출, 병당원가, 병당마진) — 2023~2026 연도 비교 가능",
    "paengi_stages": "팽이 배치 전 공정 한 줄 표 (bid, 접종일, 입상일, 권지일, 수확시작, 종균코드/종균pH/종균세대/액배양일, 살균일/살균솥/살균F0/살균2차유지분/살균중단, 접종손실률%, 배양일수/배양temp/배양hum/배양co2/긁기불량률%, 발이temp/발이hum/발이co2/억제12_14/권지불량률%, 생육temp/생육hum/생육co2, 병당중량_표시, QC, 완결, season, month)",
    "retort_vs_scrape": "살균일별 살균 조건과 다음날 접종분의 균긁기 불량률 (살균일, 접종일, 품목, 솥수, F0최소, 이차유지최소, 중단, 배양출고, 긁기불량률%)",
    "tank_batches": "액탱크 배치 종합 (spawn_tanks + 접종 시 pH + 살균솥 액탱크 사이클 살균솥/살균2차유지분/살균F0 + 배양출고/긁기불량률% + 병당중량/QC)",
    "retort_cycles": "살균솥 살균 사이클 (unit 솥 '1호기', start 시작, end 냉각완료, 승온분, 1차유지분, 1차유지온도, 2차승온분, 2차유지분, 2차유지온도, 최고온도, 최고시각, 냉각분, 총가동분, F0, 최고압력 kPa, 중단횟수, 완료, 구분 배지/액탱크/이상)",
    "retort_daily": "살균솥 기록 일별 요약 (unit 솥, tag 온도/압력/스팀/상부/하부/설정 등, date, min, max, mean, n, na 결측, values 설정값)",
    "repairs": "설비 수리 기록 — 기계수리.xlsx(구역별 시트 + 일백 냉동 공사) · 가습기 수리내역 (date, 구역, 장소 방·장비 쉼표로 여러 개, "
               "방 환경 기록 방 이름 쉼표로 여러 개 예: 발이실06·생육실03·냉각1실·배양실1·C05·1호기, 분류 냉매·누수/팽창변·전자변/콤프/패킹/쿨러·팬·히터/구동부 등, "
               "내용, 사양, 업체 일백·디엠에스 등, 비고 합친 다른 시트의 글, 시트, 행). 특정 방의 수리는 방 열에 그 이름이 들어 있는 행",
    "repair_costs": "수리 완료 보고서 (date 수리일, 완료일, 구역, 수리명, 업체, 수리비 원, 발주처, 인원, 기간, 내용) — 일백 냉동 공사는 한두 달 치를 한 장으로 묶음",
}


def make_loaders(A, etl_logs, env_loader, paengi_loader, ss_loader, monthly_loader, anomalies_loader):
    """app.py 의 캐시 로더를 받아 표 로더 딕셔너리를 만든다 (필요할 때만 계산)."""
    ENV_COLS = ["farm", "room", "date", "crop", "run", "pause", "temp", "temp_sd", "temp_min", "temp_max", "hum", "hum_min", "hum_max", "co2", "co2_min", "co2_max",
                "sp_temp", "sp_hum", "sp_co2", "cool", "heat", "humid", "dehum", "err", "track_err", "day_first", "day_last", "reset_n", "sensor_bad_n", "err_n", "mode_last"]
    L = {
        "env": lambda: (lambda e: e[[c for c in ENV_COLS if c in e.columns]])(env_loader()),
        "paengi_batches": lambda: paengi_loader()[0],
        "paengi_batch_env": lambda: paengi_loader()[1],
        "ss_batches": lambda: ss_loader()[0],
        "ss_batch_env": lambda: ss_loader()[1],
        "paengi_monthly": lambda: monthly_loader()[0],
        "ss_monthly": lambda: monthly_loader()[1],
        "cult_monthly": lambda: monthly_loader()[2],
        "anomalies": lambda: anomalies_loader(60),
        "tank_rank": lambda: A.tank_rank(),
        "cost_monthly": lambda: A.cost_monthly()[0],
        "stock_recon_paengi": lambda: A.stock_recon("팽이")[0],
        "stock_recon_ss": lambda: A.stock_recon("새송이")[0],
        "retort_daily": lambda: __import__("mushroom_app.etl_retort", fromlist=["load_daily"]).load_daily(),
        "retort_cycles": lambda: __import__("mushroom_app.etl_retort", fromlist=["load_cycles"]).load_cycles(),
        "tank_batches": lambda: A.tank_batches(),
        "paengi_stages": lambda: A.paengi_stage_table(),
        "loads": lambda: A.loads_with_cycles(),
        "sales": lambda: A.sales_monthly(),
        "retort_vs_scrape": lambda: A.retort_vs_scrape(),
        "repairs": lambda: A.repairs(),
        "repair_costs": lambda: A.repair_costs(),
    }
    for k in ("cult_scrape_paengi", "cult_scrape_ss", "cult_kwonji", "cult_inoc", "paengi_harvest", "paengi_qc", "ss_daily", "ss_room_pack",
              "materials", "electric", "fuel", "bottling", "ledger", "spawn_tanks", "tank_daily"):
        L[k] = (lambda kk=k: etl_logs.load(kk))
    return L


def _describe(T):
    out = []
    for k in T.keys():
        try:
            df = T[k]
        except Exception as e:
            out.append(f"- {k}: (읽기 실패 {e})"); continue
        if df is None or len(df) == 0:
            out.append(f"- {k}: 0행 — {TABLE_DOC.get(k, '')}"); continue
        span = ""
        for dc in ("date", "입상일", "긁기일", "권지일", "일자", "월", "m"):
            if dc in df.columns:
                s = pd.to_datetime(df[dc], errors="coerce") if dc not in ("월", "m") else df[dc].astype(str)
                span = f", 기간 {s.min()}~{s.max()}"
                if dc not in ("월", "m"):
                    span = f", 기간 {s.min().date()}~{s.max().date()}"
                break
        cols = ", ".join(map(str, df.columns[:40]))
        out.append(f"- {k}: {len(df):,}행{span}\n    설명: {TABLE_DOC.get(k, '')}\n    열: {cols}")
    return "\n".join(out)


def _run_pandas(code, T):
    if _FORBID.search(code):
        return "금지된 구문이 있어 실행하지 않았습니다(import/파일/OS 접근 불가). 표 조회·집계만 하세요."
    import ast
    buf = io.StringIO()
    ns = {"pd": pd, "np": np, "T": T}
    try:
        with contextlib.redirect_stdout(buf), pd.option_context("display.width", 200, "display.max_columns", 40, "display.max_rows", 60):
            tree = ast.parse(code)
            # 마지막 문장이 식이면 그 값을 출력한다 (노트북처럼)
            if tree.body and isinstance(tree.body[-1], ast.Expr):
                last = ast.Expression(tree.body.pop().value)
                exec(compile(tree, "<ai>", "exec"), ns)
                val = eval(compile(last, "<ai>", "eval"), ns)
                if val is not None:
                    print(val)
            else:
                exec(compile(tree, "<ai>", "exec"), ns)
    except Exception:
        return "오류:\n" + traceback.format_exc(limit=2)
    s = buf.getvalue()
    return s[:MAX_OUT] + ("\n...(잘림)" if len(s) > MAX_OUT else "") if s.strip() else "(출력 없음 — print() 나 마지막 줄에 값을 두세요)"


def _find_files(keywords, limit=30):
    kws = [k.lower() for k in str(keywords).split() if k.strip()]
    hits = []
    for root in _roots():
        for dp, dn, fn in os.walk(root):
            dn[:] = [d for d in dn if not d.startswith(("#", "$", ".")) and "백업" not in d and "app_cache" not in d]
            for f in fn:
                if f.startswith("~$"):
                    continue
                low = f.lower()
                if all(k in low for k in kws):
                    p = os.path.join(dp, f)
                    try:
                        st_ = os.stat(p)
                        hits.append((st_.st_mtime, p, st_.st_size))
                    except OSError:
                        pass
            if len(hits) > 400:
                break
    hits.sort(reverse=True)
    if not hits:
        return "일치하는 파일 없음"
    return "\n".join(f"{p}  ({sz/1e6:.1f}MB, {pd.Timestamp(mt, unit='s'):%Y-%m-%d})" for mt, p, sz in hits[:limit]) + (f"\n... 외 {len(hits)-limit}개" if len(hits) > limit else "")


def _allowed(path):
    ap = os.path.abspath(path)
    return any(ap.startswith(r) for r in _roots())


def _read_excel(path, sheet=None, max_rows=40, skip_rows=0):
    if not _allowed(path):
        return "허용된 폴더 밖의 파일은 읽을 수 없습니다."
    if not os.path.exists(path):
        return "파일이 없습니다: " + path
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".csv", ".txt"):
            with open(path, encoding="utf-8-sig", errors="replace") as f:
                lines = [next(f) for _ in range(min(max_rows, 200))]
            return "".join(lines)
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        ws = wb[sheet] if sheet and sheet in names else wb[names[0]]
        rows = []
        for i, r in enumerate(ws.iter_rows(min_row=1 + skip_rows, max_row=skip_rows + max_rows, values_only=True)):
            if any(c is not None for c in r):
                rows.append(f"{i+1+skip_rows}: " + " | ".join("" if c is None else (c.strftime("%Y-%m-%d") if hasattr(c, "strftime") else str(c))[:24] for c in r[:30]))
        wb.close()
        head = f"시트 목록: {names}\n[시트 {ws.title}] 크기 {ws.max_row}행×{ws.max_column}열\n"
        s = head + "\n".join(rows)
        return s[:MAX_OUT] + ("\n...(잘림)" if len(s) > MAX_OUT else "")
    except Exception as e:
        return f"읽기 실패: {e}"


SYSTEM = f"""당신은 '{C.APP_NAME}' 앱 안의 데이터 도우미입니다. 버섯(팽이·새송이) 재배 농장의 재배 관리자·직원이 한국어로 묻습니다.
농장 구성: 진성팽이재배사(1433: 발이실01~18, 생육실01~12, 배양실1·2, 냉각실), 연우1번재배사(1377: A01~A10, B01~B10), 연우2번재배사(1378: C01~C10, D01~D14). C02는 창고, B10·D09~D14는 저온실.
팽이 흐름: 접종 → 배양(약 28일) → 균긁기=발이실 입상(입상일=1일차) → 억제 → 권지=생육실 입실(약 16~17일차) → 수확(약 27~28일차). 새송이: 입상(0일차) → 7일 뒤집기 → 11일 솎기 → 13~14일 수확 → 폐상 → 재입상(약 17~18일). 2026-04-09 새송이 관리자 교체.
답변 규칙:
- 데이터 질문은 반드시 도구로 실제 수치를 조회한 뒤 답하고, 어떤 표·기간을 썼는지 한 줄로 밝힙니다. 추측으로 숫자를 만들지 않습니다.
- 먼저 describe_tables 로 표를 확인하고 run_pandas 로 집계합니다. 파일을 찾는 질문은 find_files → 필요하면 read_excel.
- 숫자는 단위와 함께 간결하게. 표가 필요하면 마크다운 표. 길게 늘어놓지 말고 핵심 먼저.
- 팽이 병당중량은 '병당중량_표시'(단일방 값 없으면 같은 날 합산 평균) 를 씁니다. QC 서식 A(6월까지)와 B(7월~)는 점수 체계가 달라 직접 비교하지 않습니다.
- 새송이 배치별 수확량·병당g는 2025-10~2026-04(방별 포장량)와 2026-03~현재(정규철 실별QC) 두 출처가 합쳐져 있습니다.
  실별QC에는 방 단위 QC 점수(생육·대모양·갓상태·병해·생산량 합계, QC종합)와 현장 메모도 있습니다. 그 사이 기간만 비어 있습니다.
- 배치 표에는 바깥 날씨도 붙어 있습니다: 팽이는 배양외기·발이외기·생육외기(구간 평균 기온)와 발이비일수·생육비일수,
  새송이는 외기온·비일수. 계절 효과를 볼 때 쓰되, 외기온은 계절의 다른 이름이라 인과로 단정하지 않습니다.
- 팽이 발이실·생육실은 습도를 목표값(%)으로 제어하지 않고 가습기를 정해진 주기로 돌립니다(env 표의 가습주기, 예 "50초/15초").
  그 방들은 sp_hum(설정습도)이 비어 있는 게 정상입니다.
- 모르는 것, 자료에 없는 것은 없다고 말합니다.
"""


def system_for(ctx=None, role="관리자"):
    """지금 보고 있는 화면을 지시문 끝에 붙인다. 등급이 직원이면 볼 수 없는 자료도 알려준다."""
    out = SYSTEM
    if ctx:
        out += "\n지금 사용자가 보고 있는 화면: " + ctx + "\n질문에 '이 배치'·'이 방'·'여기' 처럼 가리키는 말이 나오면 이 화면을 뜻합니다."
    if role != "관리자":
        out += ("\n이 사용자는 '직원' 등급이라 원가·매출·재고 대사·수리비 자료를 볼 수 없습니다. "
                "그 쪽 질문에는 자료를 찾지 말고 '권한이 없어 알려드릴 수 없습니다. 관리자에게 문의하세요' 라고만 답합니다.")
    return out


HIDE_FOR_STAFF = ("cost_monthly", "sales", "stock_recon_paengi", "stock_recon_ss",
                  "ledger", "materials", "electric", "fuel", "repair_costs")


def loaders_for(loaders, role):
    """등급에 맞는 표만 남긴다."""
    if role == "관리자":
        return loaders
    return {k: v for k, v in loaders.items() if k not in HIDE_FOR_STAFF}


def ask(messages, loaders, on_tool=None, ctx=None, role="관리자"):
    """messages: Anthropic 형식 대화. 도구 호출을 처리하며 최종 답 텍스트와 갱신된 messages 를 돌려준다."""
    import anthropic
    client = anthropic.Anthropic(api_key=api_key())
    loaders = loaders_for(loaders, role)
    T = _tables(loaders)
    msgs = list(messages)
    for _ in range(MAX_TURNS):
        resp = client.messages.create(model=MODEL, max_tokens=2000, system=system_for(ctx, role), tools=TOOLS, messages=msgs)
        msgs.append({"role": "assistant", "content": [b.model_dump() for b in resp.content]})
        if resp.stop_reason != "tool_use":
            text = "".join(b.text for b in resp.content if b.type == "text")
            return text, msgs
        results = []
        for b in resp.content:
            if b.type != "tool_use":
                continue
            a = b.input or {}
            if on_tool:
                on_tool(b.name, a)
            try:
                if b.name == "describe_tables":
                    out = _describe(T)
                elif b.name == "run_pandas":
                    out = _run_pandas(a.get("code", ""), T)
                elif b.name == "find_files":
                    out = _find_files(a.get("keywords", ""), int(a.get("limit") or 30))
                elif b.name == "read_excel":
                    out = _read_excel(a.get("path", ""), a.get("sheet"), int(a.get("max_rows") or 40), int(a.get("skip_rows") or 0))
                else:
                    out = "알 수 없는 도구"
            except Exception as e:
                out = f"도구 오류: {e}"
            results.append({"type": "tool_result", "tool_use_id": b.id, "content": str(out)})
        msgs.append({"role": "user", "content": results})
    return "도구 호출이 너무 많아 중단했습니다. 질문을 좁혀 다시 물어봐 주세요.", msgs
